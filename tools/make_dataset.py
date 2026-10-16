import argparse
import csv
import datetime as dt
import json
import math
import random
import string
import sys
from collections import Counter
from fractions import Fraction as F
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")


class Retry(Exception):
    pass


def num(rng, lo, hi, d=1):
    k = 10 ** d
    return F(rng.randint(math.ceil(lo * k), max(math.ceil(lo * k), math.floor(hi * k))), k)


def rnd(x, d=0, check=True):
    x = F(x)
    m = abs(x) * 10 ** d
    whole, rem = divmod(m.numerator, m.denominator)
    frac = F(rem, m.denominator)
    if check and frac and abs(frac - F(1, 2)) < F(1, 1000):
        raise Retry("sát ranh giới làm tròn")
    n = whole + (frac >= F(1, 2))
    s = str(n).rjust(d + 1, "0")
    if d:
        s = s[:-d] + "." + s[-d:]
    return "-" + s if x < 0 and n else s


def q(x, d=1):
    return F(rnd(x, d, check=False))


def short(x):
    s = rnd(x, 1, check=False)
    return s[:-2] if s.endswith(".0") else s


def top(items, key, low=False):
    s = sorted(items, key=key, reverse=not low)
    if len(s) < 2 or key(s[0]) == key(s[1]):
        raise Retry("hoà")
    return s[0]


def split(rng, total, n, slack=0.0):
    rest = total - rng.randint(0, int(total * slack))
    cuts = sorted(rng.sample(range(1, rest), n - 1))
    return [b - a for a, b in zip([0] + cuts, cuts + [rest])]


def and_list(xs, word="and"):
    return ", ".join(xs[:-1]) + f" {word} " + xs[-1]


def money(v, dollar=False):
    s = "—" if v == 0 else f"{v:,}"
    return "$" + s if dollar else s


def tree_children(items):
    kids, stack = {}, []
    for lvl, name in items:
        while stack and stack[-1][0] >= lvl:
            stack.pop()
        if stack:
            kids.setdefault(stack[-1][1], []).append(name)
        stack.append((lvl, name))
    return kids


def prune(rng, items, p):
    out, skip = [], None
    for lvl, name in items:
        if skip is not None and lvl > skip:
            continue
        skip = None
        if lvl >= 2 and rng.random() < p:
            skip = lvl
            continue
        out.append((lvl, name))
    return out


class Sheet:
    def __init__(self, ws):
        self.ws, self.r = ws, 0

    def row(self, vals=()):
        self.r += 1
        for j, v in enumerate(vals):
            if isinstance(v, F):
                v = int(v) if v.denominator == 1 else float(v)
            if v is not None:
                self.ws.cell(self.r, j + 1, v)
        return self.r

    def merge(self, r1, c1, r2, c2):
        self.ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def workbook(title):
    wb = openpyxl.Workbook()
    wb.active.title = title
    return wb, Sheet(wb.active)


ACT = [(0, "Total, all activities"), (1, "Personal care activities"), (2, "Sleeping"), (2, "Grooming"),
       (1, "Eating and drinking"), (1, "Household activities"), (2, "Housework"), (2, "Food preparation and cleanup"),
       (2, "Lawn and garden care"), (2, "Household management"), (1, "Purchasing goods and services"),
       (2, "Consumer goods purchases"), (1, "Caring for and helping household members"),
       (2, "Caring for and helping household children"), (2, "Caring for household adults"),
       (1, "Working and work-related activities"), (2, "Working"), (1, "Educational activities"),
       (1, "Organizational, civic, and religious activities"), (1, "Leisure and sports"),
       (2, "Socializing and communicating"), (2, "Watching TV"),
       (2, "Participating in sports, exercise, and recreation"), (2, "Relaxing and thinking"),
       (1, "Telephone calls, mail, and e-mail"), (1, "Other activities, not elsewhere classified")]
ACT_BASE = {"Sleeping": 8.4, "Grooming": .7, "Eating and drinking": 1.1, "Housework": .8,
            "Food preparation and cleanup": .8, "Lawn and garden care": .2, "Household management": .2,
            "Consumer goods purchases": .4, "Caring for and helping household children": 1.2,
            "Caring for household adults": .06, "Educational activities": .1,
            "Organizational, civic, and religious activities": .3, "Socializing and communicating": .6,
            "Watching TV": 2.0, "Participating in sports, exercise, and recreation": .3, "Relaxing and thinking": .2,
            "Telephone calls, mail, and e-mail": .2, "Other activities, not elsewhere classified": .2}
ACT_GROUPS = [("age of youngest own household child",
               ["Own household children under 18", "Own household children under 6",
                "Own household children 6-17, none younger"]),
              ("age of respondent", ["Ages 25 to 34", "Ages 35 to 44", "Ages 45 to 54"]),
              ("number of own household children",
               ["One own household child", "Two own household children", "Three or more own household children"])]
EMP = {"Employed full time": "full-time employed", "Employed part time": "part-time employed",
       "Not employed": "not-employed"}


def gen_activity(rng):
    pop_a, pop_b = rng.choice([("Married mothers", "Married fathers"), ("Single mothers", "Single fathers"),
                               ("Married women", "Married men"), ("Unmarried women", "Unmarried men")])
    grouping, groups = rng.choice(ACT_GROUPS)
    st_a = ["Employed full time", "Employed part time", "Not employed"]
    st_b = st_a if rng.random() < .4 else ["Employed full time", "Not employed"]
    cols = [(pop_a, s) for s in st_a] + [(pop_b, s) for s in st_b]
    items = prune(rng, ACT, .15)
    kids = tree_children(items)
    work = {"Employed full time": 5.6, "Employed part time": 3.3, "Not employed": .15}
    hours, pct = {}, {}
    for g in groups:
        for c in cols:
            for lvl, a in reversed(items):
                if lvl == 0:
                    h, p = F(24), F(100)
                elif a in kids:
                    h = sum(hours[g, c, k] for k in kids[a]) + num(rng, 0, .3, 2)
                    p = min(F(100), max(pct[g, c, k] for k in kids[a]) + num(rng, 0, 8, 1))
                else:
                    base = work[c[1]] if a == "Working" else ACT_BASE.get(a, .3)
                    h = num(rng, base * .6, base * 1.4, 2)
                    p = num(rng, 99, 100, 1) if a == "Sleeping" else num(rng, 3, 96, 1)
                hours[g, c, a], pct[g, c, a] = h, p

    n = rng.randint(1, 9)
    y = rng.randint(2005, 2018)
    wb, sh = workbook(f"Table A-{n}")
    nc = 1 + 2 * len(cols)
    r = sh.row([f"Table A-{n}. Time spent in primary activities 1 and the percent of {pop_a.lower()} and "
                f"{pop_b.lower()} who did the activities on an average day by employment status and {grouping}, "
                f"average for the combined years {y}-{(y + 4) % 100:02d}"])
    sh.merge(r, 1, r, nc)
    r = sh.row(["Activity", "Average hours per day spent in primary activities 1 "] + [None] * (len(cols) - 1)
               + ["Percent doing the activity on an average day"])
    sh.merge(r, 1, r + 2, 1)
    sh.merge(r, 2, r, 1 + len(cols))
    sh.merge(r, 2 + len(cols), r, nc)
    block = [pop_a] + [None] * (len(st_a) - 1) + [pop_b + " 2"] + [None] * (len(st_b) - 1)
    r = sh.row([None] + block * 2)
    for b in (2, 2 + len(cols)):
        sh.merge(r, b, r, b + len(st_a) - 1)
        sh.merge(r, b + len(st_a), r, b + len(cols) - 1)
    sh.row([None] + [s for _, s in cols] * 2)
    pad = {0: "", 1: "  ", 2: "      "}
    for g in groups:
        sh.row([g])
        sh.row([None])
        for lvl, a in items:
            sh.row([pad[lvl] + a + (" " if lvl == 0 else "")] + [hours[g, c, a] for c in cols]
                   + [pct[g, c, a] for c in cols])
        sh.row([None])
    sh.row(["1 Primary activities are those that respondents identified as their main activity. "
            "Other activities done simultaneously are not included."])
    sh.row([f"2 Estimates for {pop_b.lower()} employed part time are not shown separately." if len(st_b) == 2
            else f"2 Includes {pop_b.lower()} of all ages."])
    sh.row(["NOTE: Data refer to persons 18 years and over. Estimates are averages for the combined years."])

    def desc(c):
        return f"{EMP[c[1]]} {c[0].lower()}"

    def tag(c):
        return f"{EMP[c[1]]} {c[0].split()[-1].lower()}"

    leaves = [a for lvl, a in items if lvl > 0 and a not in kids]

    def q_best_group():
        c, a = rng.choice(cols), rng.choice(leaves)
        g = top(groups, key=lambda g: hours[g, c, a])
        return (f"Across the three groups shown in the table (by {grouping}), for which group do {desc(c)} spend "
                f"the most average hours per day on '{a}'? Answer with only the group name exactly as it appears "
                f"in the table.", g)

    def q_compare():
        g, a = rng.choice(groups), rng.choice(leaves)
        c1, c2 = rng.choice(cols[:len(st_a)]), rng.choice(cols[len(st_a):])
        if rng.random() < .5:
            c1, c2 = c2, c1
        h1, h2 = hours[g, c1, a], hours[g, c2, a]
        if h1 == h2:
            raise Retry
        w = c1 if h1 > h2 else c2
        return (f"Compare the average hours per day spent on '{a}' by {desc(c1)} and {desc(c2)} in the '{g}' group. "
                f"Who spends more, and by how many hours? Answer as `group, difference`, where group is either "
                f"`{tag(c1)}` or `{tag(c2)}`, and the difference is in hours using exactly two decimal places.",
                f"{tag(w)}, {rnd(abs(h1 - h2), 2)}")

    def q_percent():
        g, c = rng.choice(groups), rng.choice(cols)
        a = rng.choice([a for lvl, a in items if lvl > 0])
        return (f"In the '{g}' group, what percent of {desc(c)} did '{a}' on an average day? Answer with only the "
                f"number using exactly one decimal place.", rnd(pct[g, c, a], 1))

    def q_child():
        p = rng.choice([k for k, v in kids.items() if lvl_of[k] == 1 and len(v) >= 2])
        g, c = rng.choice(groups), rng.choice(cols)
        a = top(kids[p], key=lambda k: hours[g, c, k])
        return (f"In the '{g}' group, which activity listed directly under '{p}' takes up the most average hours per "
                f"day for {desc(c)}? Answer with only the activity name exactly as it appears in the table, without "
                f"leading spaces.", a)

    def q_sum():
        g, c = rng.choice(groups), rng.choice(cols)
        a1, a2 = rng.sample(leaves, 2)
        return (f"For {desc(c)} in the '{g}' group, what is the sum of the average hours per day spent on '{a1}' and "
                f"the average hours per day spent on '{a2}'? Answer with only the number using exactly two decimal "
                f"places.",
                rnd(hours[g, c, a1] + hours[g, c, a2], 2))

    lvl_of = {a: lvl for lvl, a in items}
    return wb, [q_best_group, q_compare, q_percent, q_child, q_sum]


CPI_EXP = [(1, "Food and beverages"), (2, "Food"), (3, "Food at home"), (3, "Food away from home"),
           (2, "Alcoholic beverages"), (1, "Housing"), (2, "Shelter"), (2, "Fuels and utilities"),
           (2, "Household furnishings and operations"), (1, "Apparel"), (2, "Men's and boys' apparel"),
           (2, "Women's and girls' apparel"), (2, "Footwear"), (1, "Transportation"), (2, "Private transportation"),
           (3, "New and used motor vehicles"), (3, "Motor fuel"), (3, "Motor vehicle insurance"),
           (2, "Public transportation"), (3, "Airline fares"), (1, "Medical care"), (2, "Medical care commodities"),
           (2, "Medical care services"), (1, "Recreation"), (2, "Video and audio"),
           (2, "Pets, pet products and services"), (2, "Sporting goods"), (1, "Education and communication"),
           (2, "Education"), (2, "Communication"), (1, "Other goods and services"),
           (2, "Tobacco and smoking products"), (2, "Personal care")]
CPI_GRP = [(1, "Commodities"), (2, "Nondurables"), (2, "Durables"), (1, "Services"), (2, "Rent of shelter"),
           (2, "Transportation services"), (2, "Other services"), (1, "All items less food and energy"),
           (1, "Energy")]
MONTHS = ["Jan.", "Feb.", "Mar.", "Apr.", "May", "June", "July", "Aug.", "Sept.", "Oct.", "Nov.", "Dec."]


def gen_cpi(rng):
    Y, mi = rng.randint(2012, 2025), rng.randint(0, 11)
    M, Mp = MONTHS[mi], MONTHS[mi - 1]
    Yp = Y if mi else Y - 1
    exp = [(0, "All items")] + prune(rng, CPI_EXP, .15)
    grp = CPI_GRP
    w = {"All items": 100000}
    kids = tree_children(exp)

    def assign(name):
        ks = kids.get(name, [])
        if ks:
            for k, v in zip(ks, split(rng, w[name], len(ks), 0 if name == "All items" else .15)):
                w[k] = v
                assign(k)
    assign("All items")
    w["Commodities"], w["Services"] = split(rng, 100000, 2)
    gkids = tree_children(grp)
    for p in ("Commodities", "Services"):
        for k, v in zip(gkids[p], split(rng, w[p], len(gkids[p]), .1)):
            w[k] = v
    w["All items less food and energy"] = rng.randint(75000, 82000)
    w["Energy"] = rng.randint(6000, 8500)

    rows = exp + grp
    prev, cur, mchg, ychg = {}, {}, {}, {}
    for _, n in rows:
        prev[n] = num(rng, 110, 420, 3)
        cur[n] = q(prev[n] * (1 + num(rng, -1.5, 1.5, 1) / 100), 3)
        mchg[n] = q((cur[n] / prev[n] - 1) * 100, 1)
        ychg[n] = num(rng, -4, 9, 1)

    wb, sh = workbook("Sheet0")
    title = rng.choice([
        "Table 1C. Chained Consumer Price Index for All Urban Consumers (C-CPI-U): U.S. city average, "
        "by expenditure category and commodity and service group",
        "Table 3. Consumer Price Index for Urban Wage Earners and Clerical Workers (CPI-W): U.S. city average, "
        "by expenditure category and commodity and service group",
        f"Table 2. Consumer Price Index for All Urban Consumers (CPI-U): {rng.choice(['Northeast', 'Midwest', 'South', 'West'])} "
        "region average, by expenditure category and commodity and service group"])
    r = sh.row(["", title])
    sh.merge(r, 2, r, 7)
    r = sh.row(["", rng.choice(["[December 1999=100, unless otherwise noted]", "[1982-84=100, unless otherwise noted]"])])
    sh.merge(r, 2, r, 7)
    sh.row(["", ""])
    r = sh.row(["Indent Level", "Item and group", f"Relative importance {Y - 1}", "Unadjusted indexes", None,
                f"Unadjusted\npercent change to\n{M} {Y} from--", None])
    sh.merge(r, 4, r, 5)
    sh.merge(r, 6, r, 7)
    sh.row([None, None, None, f"{Mp}\n{Yp}", f"{M}\n{Y}", f"{M}\n{Y - 1}", f"{Mp}\n{Yp}"])
    sh.row(["", ""])
    sh.row([None])
    for header, part in (("Expenditure category", exp), ("Commodity and service group", grp)):
        sh.row([0, header])
        for lvl, n in part:
            sh.row([lvl, n, F(w[n], 1000), prev[n], cur[n], ychg[n], mchg[n]])
    sh.row([None])
    sh.row(["", "Footnotes"])
    sh.row(["", "(1) Not seasonally adjusted. Relative importance is based on the previous year's expenditure weights."])

    names = [n for _, n in rows]
    level1 = [n for lvl, n in exp if lvl == 1]
    allkids = {**kids, **gkids}
    period = f"from {M} {Y - 1} to {M} {Y}"

    def q_pp():
        a, b = rng.sample(names, 2)
        if ychg[a] == ychg[b]:
            raise Retry
        if ychg[a] < ychg[b]:
            a, b = b, a
        return (f"How many percentage points higher was the unadjusted percent change {period} for {a} than for {b}? "
                f"Answer with only the number using exactly one decimal place.", rnd(ychg[a] - ychg[b], 1))

    def q_weight():
        a, b = rng.sample(level1, 2)
        if w[a] < w[b]:
            a, b = b, a
        return (f"How many percentage points larger is the relative importance of {a} than that of {b}? "
                f"Answer with only the number using exactly three decimal places.", rnd(F(w[a] - w[b], 1000), 3))

    def q_rank():
        p = rng.choice([k for k, v in allkids.items() if k != "All items" and len(v) >= 3
                        and not any("," in x for x in v)])
        ks = allkids[p]
        if len({ychg[k] for k in ks}) < len(ks):
            raise Retry
        order = sorted(ks, key=lambda k: -ychg[k])
        return (f"Rank the categories directly under “{p}” by their unadjusted percent change to {M} {Y} from "
                f"{M} {Y - 1}, from highest to lowest. Answer with the category names in that order, separated by a "
                f"comma and one space, without numbering.", ", ".join(order))

    def q_index():
        a = rng.choice(names)
        return (f"By how much did the unadjusted index for {a} change from {Mp} {Yp} to {M} {Y}? Answer with only "
                f"the number using exactly three decimal places, with a leading minus sign if it fell.",
                rnd(cur[a] - prev[a], 3))

    def q_level3():
        a = top([n for lvl, n in rows if lvl == 3], key=ychg.get)
        return (f"Among the items at indent level 3, which had the largest unadjusted percent change {period}? "
                f"Answer with only the item name exactly as it appears in the table.", a)

    def q_count():
        t = num(rng, 1, 5, 1)
        return (f"How many of the expenditure categories at indent level 1 had an unadjusted percent change {period} "
                f"above {rnd(t, 1)} percent? Answer with only the integer.", str(sum(ychg[k] > t for k in level1)))

    return wb, [q_pp, q_weight, q_rank, q_index, q_level3, q_count]


def gen_annual(rng):
    y0 = rng.randint(1948, 1990)
    years = list(range(y0, y0 + rng.randint(30, min(75, 2025 - y0))))
    pop, part = rng.uniform(95000, 190000), rng.uniform(.57, .64)
    ur, ag = rng.uniform(.03, .07), rng.uniform(.04, .11)
    rows = []
    for y in years:
        pop *= 1 + rng.uniform(.008, .018)
        part = min(.68, max(.55, part + rng.uniform(-.006, .008)))
        ur = min(.11, max(.03, ur + (rng.uniform(.012, .03) if rng.random() < .12 else rng.uniform(-.006, .004))))
        ag *= rng.uniform(.95, 1.02)
        P = round(pop)
        LF = round(P * part)
        U = round(LF * ur)
        E = LF - U
        A = round(E * ag)
        rows.append(dict(year=y, pop=P, lf=LF, lfp=q(F(LF * 100, P)), emp=E, epop=q(F(E * 100, P)), ag=A,
                         nonag=E - A, un=U, ur=q(F(U * 100, LF)), nilf=P - LF))

    wb, sh = workbook(f"cpsaat{rng.randint(1, 3):02d}")
    r = sh.row([f"HOUSEHOLD DATA\nANNUAL AVERAGES\n 1.  Employment status of the civilian noninstitutional "
                f"population, {y0} to date"])
    sh.merge(r, 1, r, 11)
    sh.row(["[Numbers in thousands]"])
    sh.row([None])
    r = sh.row(["Year", "Civilian\nnoninstitu-\ntional\npopulation", "Civilian labor force"] + [None] * 7
               + ["Not in\nlabor force"])
    sh.row([None, None, "Total", "Percent of\npopulation", "Employed", None, None, None, "Unemployed"])
    sh.row([None, None, None, None, "Total", "Percent of\npopulation", "Agri-\nculture",
            "Nonagri-\ncultural\nindustries", "Number", "Percent of\nlabor force"])
    for a, b, c, d in [(r, 1, r + 2, 1), (r, 2, r + 2, 2), (r, 3, r, 10), (r, 11, r + 2, 11), (r + 1, 3, r + 2, 3),
                       (r + 1, 4, r + 2, 4), (r + 1, 5, r + 1, 8), (r + 1, 9, r + 1, 10)]:
        sh.merge(a, b, c, d)
    sh.row([None])
    sh.row(["Persons 16 years of age and over"])
    for x in rows:
        sh.row([x["year"], x["pop"], x["lf"], x["lfp"], x["emp"], x["epop"], x["ag"], x["nonag"], x["un"], x["ur"],
                x["nilf"]])
    r = sh.row(["NOTE: Data for some years are not strictly comparable with data for earlier years because of "
                "population control adjustments."])
    sh.merge(r, 1, r, 11)

    def q_drop():
        diffs = [(b["year"], a["emp"] - b["emp"]) for a, b in zip(rows, rows[1:])]
        y, d = top(diffs, key=lambda t: t[1])
        if d <= 0:
            raise Retry
        return ("In which year did the number of employed people fall the most compared with the previous year, and "
                "by how many thousand? Answer as `year, decrease`, separated by a comma and one space, as integers "
                "without commas or units.", f"{y}, {d}")

    def q_best_year():
        mname, mk, low = rng.choice([("employment in agriculture", "ag", False), ("the unemployment rate", "ur", True),
                                     ("the labor force participation rate", "lfp", False),
                                     ("the number of unemployed people", "un", True)])
        cname, ck, cshort = rng.choice([("total employed population", "emp", "total employed"),
                                        ("civilian labor force", "lf", "labor force"),
                                        ("population not in the labor force", "nilf", "not in labor force")])
        x = top(rows, key=lambda r: r[mk], low=low)
        return (f"Find the year in which {mname} was the {'lowest' if low else 'highest'} and state the {cname} "
                f"(in thousands) in that year. Answer as `year, {cshort}`, separated by a comma and one space, with "
                f"the {cshort} as an integer without commas or units.", f"{x['year']}, {x[ck]}")

    def q_filter_lf():
        for _ in range(400):
            t = F(rng.randint(6, 18), 2)
            lim = round(rng.choice(rows)["lf"], -3)
            m = [x for x in rows if x["ur"] > t and x["lf"] < lim]
            if len(m) == 1:
                return (f"Identify the year when the unemployment rate exceeded {short(t)}% and the total civilian "
                        f"labor force was below {lim:,} (in thousands). Answer with only the year.", str(m[0]["year"]))
        raise Retry

    def q_filter_un():
        for _ in range(400):
            lim = round(rng.choice(rows)["un"], -2)
            z = F(rng.randint(108, 128), 2)
            m = [x for x in rows if x["un"] < lim and x["epop"] > z]
            if len(m) == 1:
                return (f"In which year was the number of unemployed individuals (in thousands) below {lim:,} while "
                        f"the employed percentage of the population was above {short(z)}%? Answer with only the year.",
                        str(m[0]["year"]))
        raise Retry

    def q_count():
        t = F(rng.randint(8, 16), 2)
        a = rng.choice(years[:-10])
        b = rng.randint(a + 9, min(a + 30, years[-1]))
        c = sum(x["ur"] >= t for x in rows if a <= x["year"] <= b)
        return (f"In how many years from {a} to {b}, inclusive, was the unemployment rate at or above {short(t)} "
                f"percent? Answer with only the integer.", str(c))

    def q_avg():
        a = rng.choice(years[:-10])
        b = a + rng.randint(4, 9)
        sel = [x["ur"] for x in rows if a <= x["year"] <= b]
        return (f"What was the average annual unemployment rate from {a} to {b}, inclusive? Answer with only the "
                f"number rounded to two decimal places.", rnd(sum(sel) / len(sel), 2))

    return wb, [q_drop, q_best_year, q_filter_lf, q_filter_un, q_count, q_avg]


D_POPS = [("Hispanic or Latino", "Hispanic or Latino ethnicity", "detailed ethnic group",
           ["Mexican", "Puerto Rican", "Cuban", "Central American", "South American", "Other Hispanic or Latino"]),
          ("Asian", "Asian", "detailed Asian group",
           ["Asian Indian", "Chinese", "Filipino", "Japanese", "Korean", "Vietnamese", "Other Asian"])]
D_STATS = [("Civilian noninstitutional population", "pop"), ("Civilian labor force", "lf"),
           ("Participation rate", "lfp"), ("Employed", "emp"), ("Employment-population ratio", "epop"),
           ("Unemployed", "un"), ("Unemployment rate", "ur"), ("Not in labor force", "nilf")]


def rates(pop, lf, un):
    emp = lf - un
    return dict(pop=pop, lf=lf, emp=emp, un=un, nilf=pop - lf, lfp=q(F(lf * 100, pop)), epop=q(F(emp * 100, pop)),
                ur=q(F(un * 100, lf)))


def gen_ethnic(rng):
    pname, ptot, dlabel, all_det = rng.choice(D_POPS)
    chosen = set(rng.sample(all_det, rng.randint(3, 5)))
    det = [d for d in all_det if d in chosen]
    y = rng.randint(2015, 2024)
    yrs = [y - 1, y]
    raw = {}
    for g in det + ["_other"]:
        size = rng.uniform(300, 14000)
        for sex in ("Men", "Women"):
            p0, part, ur = size * rng.uniform(.9, 1.1), rng.uniform(.55, .8), rng.uniform(.025, .09)
            for k, yr in enumerate(yrs):
                pop = round(p0 * (1 + k * rng.uniform(-.01, .04)))
                lf = round(pop * (part + k * rng.uniform(-.015, .015)))
                un = round(lf * (ur + k * rng.uniform(-.012, .012)))
                raw[g, sex, yr] = (pop, lf, un)
    data = {}
    for yr in yrs:
        for sex in ("Men", "Women"):
            for g in det:
                data[g, sex, yr] = rates(*raw[g, sex, yr])
            tot = [sum(raw[g, sex, yr][i] for g in det + ["_other"]) for i in range(3)]
            data["Total(1)", sex, yr] = rates(*tot)
        for g in det + ["Total(1)"]:
            data[g, "Total", yr] = rates(*[data[g, "Men", yr][k] + data[g, "Women", yr][k] for k in ("pop", "lf", "un")])
    groups = ["Total(1)"] + det

    wb, sh = workbook(f"cpsaat{rng.randint(4, 7):02d}")
    nc = 1 + 2 * len(groups)
    r = sh.row([f"HOUSEHOLD DATA\nANNUAL AVERAGES \n6. Employment status of the {pname} population by sex, age, "
                f"and {dlabel}"])
    sh.merge(r, 1, r, nc)
    sh.row(["[Numbers in thousands]"])
    sh.row([None])
    r = sh.row(["Employment status, sex, and age", ptot])
    sh.merge(r, 1, r + 2, 1)
    sh.merge(r, 2, r, nc)
    r = sh.row([None] + sum(([g, None] for g in groups), []))
    for i in range(len(groups)):
        sh.merge(r, 2 + 2 * i, r, 3 + 2 * i)
    sh.row([None] + yrs * len(groups))
    sh.row([None])
    for label, sex in (("TOTAL", "Total"), ("Men, 16 years and over", "Men"), ("Women, 16 years and over", "Women")):
        sh.row([label])
        for slabel, k in D_STATS:
            sh.row([slabel] + [data[g, sex, yr][k] for g in groups for yr in yrs])
        sh.row([None])
    r = sh.row([f"(1) Includes persons in other {dlabel}s not shown separately."])
    sh.merge(r, 1, r, nc)
    r = sh.row(["NOTE: Estimates for the above groups will not sum to totals because data are not presented for "
                "all groups."])
    sh.merge(r, 1, r, nc)

    low = {"Men": "men", "Women": "women"}

    def who(g):
        return pname if g == "Total(1)" else g

    def q_increase():
        sex = rng.choice(["Men", "Women"])
        cand = rng.sample(det, 3)
        cand.sort(key=det.index)
        inc = {g: data[g, sex, y]["un"] - data[g, sex, y - 1]["un"] for g in cand}
        g = top(cand, key=inc.get)
        if inc[g] <= 0:
            raise Retry
        return (f"Among {and_list(cand)} {low[sex]} aged 16 years and over, which group had the largest increase in "
                f"the number of unemployed people from {y - 1} to {y}, and by how many thousand? Answer as "
                f"`group, increase`, separated by a comma and one space, with the increase as an integer.",
                f"{g}, {inc[g]}")

    def q_lfp_gap():
        g, yr = rng.choice(groups), rng.choice(yrs)
        d = data[g, "Men", yr]["lfp"] - data[g, "Women", yr]["lfp"]
        if d <= 0:
            raise Retry
        return (f"For {who(g)} people aged 16 years and over, how many percentage points higher was the labor force "
                f"participation rate of men than that of women in {yr}? Answer with only the number using exactly "
                f"one decimal place.", rnd(d, 1))

    def q_lf_change():
        g, sex = rng.choice(groups), rng.choice(["Men", "Women"])
        d = data[g, sex, y]["lf"] - data[g, sex, y - 1]["lf"]
        return (f"By how many thousand did the civilian labor force of {who(g)} {low[sex]} change from {y - 1} to "
                f"{y}? Answer with only the integer, with a minus sign if it decreased.", str(d))

    def q_top_ur():
        sex, yr = rng.choice(["Men", "Women"]), rng.choice(yrs)
        g = top(det, key=lambda g: data[g, sex, yr]["ur"])
        return (f"In {yr}, which {dlabel} had the highest unemployment rate among {low[sex]} aged 16 years and over? "
                f"Answer with only the group name exactly as it appears in the table.", g)

    def q_sum_un():
        g1, g2 = rng.sample(det, 2)
        sex, yr = rng.choice(["Men", "Women"]), rng.choice(yrs)
        return (f"What was the combined number of unemployed {g1} and {g2} {low[sex]} in {yr}, in thousands? "
                f"Answer with only the integer.", str(data[g1, sex, yr]["un"] + data[g2, sex, yr]["un"]))

    return wb, [q_increase, q_lfp_gap, q_lf_change, q_top_ur, q_sum_un]


OCC = [(0, "Management, professional, and related occupations"),
       (1, "Management, business, and financial operations occupations"), (1, "Professional and related occupations"),
       (0, "Service occupations"), (0, "Sales and office occupations"), (1, "Sales and related occupations"),
       (1, "Office and administrative support occupations"),
       (0, "Natural resources, construction, and maintenance occupations"),
       (1, "Farming, fishing, and forestry occupations"), (1, "Construction and extraction occupations"),
       (1, "Installation, maintenance, and repair occupations"),
       (0, "Production, transportation, and material moving occupations"), (1, "Production occupations"),
       (1, "Transportation and material moving occupations")]
RACES = ["Total", "White", "Black or African American", "Asian", "Hispanic or Latino ethnicity"]
RACE_ANS = {"White": "White", "Black or African American": "Black or African American", "Asian": "Asian",
            "Hispanic or Latino ethnicity": "Hispanic or Latino"}


def gen_occupation(rng):
    sex = rng.choice(["Women", "Men", "Both sexes"])
    sexw = {"Women": "women", "Men": "men", "Both sexes": "people"}[sex]
    y = rng.randint(2015, 2024)
    yrs = [y - 1, y]
    kids = tree_children([(-1, "_")] + OCC)
    majors = kids["_"]
    share, count = {}, {}
    for g in RACES:
        base = [rng.randint(300, 520), rng.randint(120, 260), rng.randint(150, 260), rng.randint(40, 140)]
        size = rng.randint(8000, 160000)
        for yr in yrs:
            parts = [max(40, b + rng.randint(-25, 25)) for b in base]
            last = 1000 - sum(parts)
            if last < 40:
                raise Retry
            for m, v in zip(majors, parts + [last]):
                share[g, yr, m] = F(v, 10)
                if m in kids:
                    for k, c in zip(kids[m], split(rng, v, len(kids[m]))):
                        share[g, yr, k] = F(c, 10)
            count[g, yr] = round(size * rng.uniform(.98, 1.04))
    indent = rng.random() < .5

    wb, sh = workbook(f"cpsaat{rng.randint(10, 12)}")
    nc = 1 + 2 * len(RACES)
    r = sh.row([f"HOUSEHOLD DATA\nANNUAL AVERAGES\n10. Employed {'persons' if sex == 'Both sexes' else sexw} by "
                f"occupation, race, and Hispanic or Latino ethnicity"])
    sh.merge(r, 1, r, nc)
    sh.row(["[Percent distribution]"])
    sh.row([None])
    r = sh.row(["Occupation"] + sum(([g, None] for g in RACES), []))
    sh.merge(r, 1, r + 1, 1)
    for i in range(len(RACES)):
        sh.merge(r, 2 + 2 * i, r, 3 + 2 * i)
    sh.row([None] + yrs * len(RACES))
    sh.row([None])
    sh.row([sex.upper()])
    sh.row([f"Total, 16 years and over (in thousands)"] + [count[g, yr] for g in RACES for yr in yrs])
    sh.row(["Percent of total employed"] + [100] * (2 * len(RACES)))
    for lvl, o in OCC:
        sh.row([("  " if indent and lvl else "") + o] + [share[g, yr, o] for g in RACES for yr in yrs])
    r = sh.row(["NOTE: Estimates for the above race groups will not sum to totals because data are not presented "
                "for all races. Persons whose ethnicity is identified as Hispanic or Latino may be of any race."])
    sh.merge(r, 1, r, nc)

    occs = [o for _, o in OCC]

    def lower(o):
        return o[0].lower() + o[1:]

    def q_top_race():
        o, yr = rng.choice(occs), rng.choice(yrs)
        g = top(RACES[1:], key=lambda g: share[g, yr, o])
        return (f"In {yr}, among employed {sexw} who are White, Black or African American, Asian, or Hispanic or "
                f"Latino, which group had the highest percentage working in {lower(o)}? Answer with only the group "
                f"name: White, Black or African American, Asian, or Hispanic or Latino.", RACE_ANS[g])

    def q_absdiff():
        o = rng.choice(occs)
        d = abs(share["Total", y, o] - share["Total", y - 1, o])
        if d == 0:
            raise Retry
        return (f"For employed {sexw} of all races combined, what is the absolute difference between {y - 1} and {y} "
                f"in the share working in {lower(o)}? Answer with only the number using exactly two decimal places.",
                rnd(d, 2))

    def q_count():
        g, o, yr = rng.choice(RACES[1:]), rng.choice(occs), rng.choice(yrs)
        return (f"Based on the percent distribution, approximately how many thousand employed {RACE_ANS[g]} {sexw} "
                f"worked in {lower(o)} in {yr}? Answer with only the integer, rounded to the nearest thousand.",
                rnd(count[g, yr] * share[g, yr, o] / 100, 0))

    def q_major():
        g, yr = rng.choice(RACES[1:]), rng.choice(yrs)
        o = top(majors, key=lambda o: share[g, yr, o])
        return (f"In {yr}, which major occupation group accounted for the largest share of employed {RACE_ANS[g]} "
                f"{sexw}? Answer with only the occupation group name exactly as it appears in the table.", o)

    return wb, [q_top_race, q_absdiff, q_count, q_major]


QI = [("Crop production", "111000"), ("Animal production and aquaculture", "112000"), ("Forestry and logging", "113000"),
      ("Oil and gas extraction", "211000"), ("Utilities", "221000"), ("Construction of buildings", "236000"),
      ("Heavy and civil engineering construction", "237000"), ("Specialty trade contractors", "238000"),
      ("Food manufacturing", "311000"), ("Beverage manufacturing", "312100"), ("Wood product manufacturing", "321000"),
      ("Paper manufacturing", "322000"), ("Plastics and rubber products manufacturing", "326000"),
      ("Primary metal manufacturing", "331000"), ("Fabricated metal product manufacturing", "332000"),
      ("Machinery manufacturing", "333000"), ("Transportation equipment manufacturing", "336000"),
      ("Furniture and related product manufacturing", "337000"), ("Merchant wholesalers, durable goods", "423000"),
      ("Building material and garden equipment and supplies dealers", "444000"),
      ("Food and beverage stores", "445000"), ("General merchandise stores", "452000"),
      ("Truck transportation", "484000"), ("Couriers and messengers", "492000"), ("Warehousing and storage", "493000"),
      ("Waste management and remediation services", "562000"), ("Hospitals", "622000"),
      ("Nursing and residential care facilities", "623000"), ("Sports teams and clubs", "711211"),
      ("Amusement parks and arcades", "713100"), ("Accommodation", "721000"),
      ("Food services and drinking places", "722000")]
SIZES = ["Total all sizes", "1 - 10", "11 - 49", "50 - 249", "250 - 999", "1,000+"]
SIZE_F = {"Total all sizes": 1, "1 - 10": .5, "11 - 49": 1.0, "50 - 249": 1.3, "250 - 999": 1.2, "1,000+": 1.0}


def gen_quartile(rng):
    year = rng.randint(2015, 2023)
    inds = sorted(rng.sample(QI, rng.randint(10, 16)), key=lambda x: x[1])
    aggs = ["All industries including state and local government3    ", "Private industry3    "]
    data = {}
    for key in aggs + [n for n, _ in inds]:
        base = rng.uniform(1, 7)
        for s in SIZES:
            if s == "1,000+" and key not in aggs and rng.random() < .25:
                data[key, s] = None
                continue
            m = max(F(1, 10), num(rng, base * SIZE_F[s] * .7, base * SIZE_F[s] * 1.3, 1))
            if s in ("Total all sizes", "1 - 10") or (s == "11 - 49" and rng.random() < .5):
                data[key, s] = (m, None, None, None)
            else:
                q1 = num(rng, 0, float(m) * .5, 1)
                med = q1 + num(rng, .2, max(.3, float(m) * .6), 1)
                q3 = med + num(rng, .5, max(.6, float(m) * 1.2), 1)
                data[key, s] = (m, q1, med, q3)

    wb, sh = workbook(f"{year % 100}QRTL1_00")
    sh.row(["TABLE Q1. Incidence rates1 of total recordable cases of nonfatal occupational injuries and illnesses "
            "by quartile"])
    sh.row([f"distribution and employment size, {year}"])
    sh.row(["Industry, NAICS code,2 and establishment employment size",
            "Average incidence rates for all establishments: ",
            "One-quarter of the establishments had a rate lower than or equal to: ",
            "One-half of the establishments had a rate lower than or equal to: ",
            "Three-quarters of the establishments had a rate lower than or equal to: "])
    sh.row([None, "(mean)", "(1st quartile)", "(median)", "(3rd quartile)"])
    disp = {n: f"{n} (NAICS {c})" for n, c in inds}
    for key in aggs + [n for n, _ in inds]:
        sh.row([disp.get(key, key)])
        for s in SIZES:
            v = data[key, s]
            cells = ["(-6-)"] * 4 if v is None else [x if x is not None else "(-8-)" for x in v]
            sh.row([s if s == "Total all sizes" else "   " + s] + cells)
    sh.row([None])
    sh.row(["1 The incidence rates represent the number of injuries and illnesses per 100 full-time workers."])
    sh.row(["2 North American Industry Classification System -- United States, 2017."])
    sh.row(["3 Excludes farms with fewer than 11 employees."])
    sh.row(["(-6-) Data too small to be displayed."])
    sh.row(["(-8-) Quartile data not available."])

    names = [n for n, _ in inds]

    def has_q(n, s):
        return data[n, s] is not None and data[n, s][1] is not None

    def q_gap():
        n = rng.choice(names)
        cand = [s for s in SIZES[1:] if has_q(n, s)]
        gap = {s: data[n, s][3] - data[n, s][2] for s in cand}
        s = top(cand, key=gap.get)
        return (f"Within the '{n}' industry, excluding the overall 'Total all sizes' row, which employment size class "
                f"has the largest gap between its 3rd quartile and its median incidence rate, and what is that gap? "
                f"Answer as `size class, gap`, separated by a comma and one space, with the gap using exactly one "
                f"decimal place.", f"{s}, {rnd(gap[s], 1)}")

    def q_top_mean():
        s = rng.choice(SIZES[1:])
        n = top([n for n in names if data[n, s] is not None], key=lambda n: data[n, s][0])
        return (f"Which industry has the highest average (mean) incidence rate among establishments in the '{s}' "
                f"employment size class? Consider only the industries identified by a NAICS code and ignore "
                f"industries without data for that class. Answer with only the industry name exactly as it appears "
                f"in the table.", disp[n])

    def q_yesno():
        n = rng.choice(names)
        s = rng.choice([s for s in SIZES[1:] if has_q(n, s)])
        v = data[n, s][3]
        t = v + rng.choice([-1, 1]) * num(rng, .1, 1.0, 1)
        if t <= 0:
            raise Retry
        return (f"Within the {disp[n]} industry, does the rate below which three-fourths of the '{s}' establishments "
                f"fall exceed {rnd(t, 1)}? Answer with only Yes or No.", "Yes" if v > t else "No")

    def q_count():
        v = num(rng, 1.5, 6, 1)
        c = sum(data[n, "Total all sizes"][0] >= v for n in names)
        return (f"How many of the industries identified by a NAICS code have an average (mean) incidence rate of at "
                f"least {rnd(v, 1)} for 'Total all sizes'? Answer with only the integer.", str(c))

    def q_iqr():
        n = rng.choice(names)
        s = rng.choice([s for s in SIZES[1:] if has_q(n, s)])
        return (f"For {disp[n]}, what is the interquartile range (3rd quartile minus 1st quartile) of the incidence "
                f"rates for the '{s}' employment size class? Answer with only the number using exactly one decimal "
                f"place.", rnd(data[n, s][3] - data[n, s][1], 1))

    return wb, [q_gap, q_top_mean, q_yesno, q_count, q_iqr]


F_SUB = [("Wood product manufacturing", "321"), ("Paper manufacturing", "322"),
         ("Plastics and rubber products manufacturing", "326"), ("Nonmetallic mineral product manufacturing", "327"),
         ("Primary metal manufacturing", "331"), ("Fabricated metal product manufacturing", "332"),
         ("Machinery manufacturing", "333"), ("Electrical equipment, appliance, and component manufacturing", "335"),
         ("Furniture and related product manufacturing", "337"), ("Miscellaneous manufacturing", "339")]
MATERIALS = ["Metal", "Plastic", "Glass", "Paper", "Wood", "Rubber", "Textile", "Ceramic", "Aluminum", "Copper",
             "Steel", "Concrete", "Leather", "Fiber", "Stone", "Brass", "Iron", "Zinc", "Foam", "Vinyl", "Cork",
             "Bamboo", "Nickel", "Tin", "Composite"]
CATEGORIES = ["fitting", "container", "fixture", "component", "hardware", "equipment", "structure", "device", "tool",
              "supply", "part", "product"]
PRODUCTS = ["valve", "pipe", "fastener", "hose", "bearing", "window", "door", "tile", "wire", "sheet", "tube",
            "spring", "pump", "cabinet", "panel", "box", "bag", "bottle", "shelf", "frame", "gear", "hinge", "clamp",
            "rod", "ring", "plate", "mesh", "handle", "lid", "crate", "drum", "sink", "bucket", "ladder", "bracket"]


def gen_naics(rng):
    year = rng.randint(2015, 2023)
    subs = sorted(rng.sample(F_SUB, rng.randint(5, 8)), key=lambda x: x[1])
    used, rows = set(), []

    def node(name, code, level, parent=None, marker=""):
        k = rng.uniform(.6, 1.5)
        away, tr = num(rng, .2, 3 * k, 1), num(rng, 0, 2.5 * k, 1)
        dj = away + tr + rng.choice([F(-1, 10), 0, 0, 0, 0, F(1, 10)])
        dj = max(dj, away, tr)
        other = num(rng, .3, 3.5 * k, 1)
        total = dj + other + rng.choice([F(-1, 10), 0, 0, 0, F(1, 10)])
        rows.append(dict(name=name, label=name + marker, code=code, level=level, parent=parent, total=total, dj=dj,
                         away=away, tr=tr, other=other))

    for name in ("All industries including private, state and local government(5)", "Private industry(5)",
                 "Goods-producing(5)"):
        node(name, None, 0)
    node("Manufacturing", "31-33", 2)
    for sname, scode in subs:
        node(sname, scode, 3, marker=rng.choice(["", "(5)"]))
        for gi in range(rng.randint(3, 6)):
            for _ in range(30):
                m, c = rng.choice(MATERIALS), rng.choice(CATEGORIES)
                gname = f"{m} {c} manufacturing"
                if gname not in used:
                    break
            else:
                continue
            used.add(gname)
            gcode = f"{scode}{gi + 1}"
            node(gname, gcode, 4, parent=scode)
            for li, p in enumerate(rng.sample(PRODUCTS, rng.randint(3, 9))):
                lname = f"{m} {p} manufacturing"
                if lname not in used:
                    used.add(lname)
                    node(lname, f"{gcode}1{li + 1}", 6, parent=gcode)

    wb, sh = workbook(str(year))
    r = sh.row([f"TABLE 1. Incidence rates(1)of nonfatal occupational injuries and illnesses by industry and case "
                f"types, {year}"])
    sh.merge(r, 1, r, 7)
    sh.row(["", ""])
    r = sh.row(["Industry(2)", "NAICS code(3)", "Total recordable cases",
                "Cases with days away from work, job restriction, or transfer", None, None, "Other recordable cases"])
    sh.row([None, None, None, "Total", "Cases with days away from work(4)",
            "Cases with days of job transfer or restriction", None])
    for a, b, c, d in [(r, 1, r + 1, 1), (r, 2, r + 1, 2), (r, 3, r + 1, 3), (r, 4, r, 6), (r, 7, r + 1, 7)]:
        sh.merge(a, b, c, d)
    sh.row(["", ""])
    for x in rows:
        sh.row([x["label"], x["code"], x["total"], x["dj"], x["away"], x["tr"], x["other"]])
    sh.row([None])
    for t in ["(1) The incidence rates represent the number of injuries and illnesses per 100 full-time workers.",
              "(2) Totals include data for industries not shown separately.",
              "(3) North American Industry Classification System -- United States, 2017.",
              "(4) Days-away-from-work cases include those that result in days away from work with or without job "
              "transfer or restriction.",
              "(5) Excludes farms with fewer than 11 employees."]:
        sh.row([t])

    leaves = [x for x in rows if x["level"] == 6]
    groups = [x for x in rows if x["level"] == 4]

    def q_count():
        t = num(rng, 2, 6, 1)
        c = sum(x["total"] > t for x in leaves)
        if not 3 <= c <= len(leaves) - 3:
            raise Retry
        return (f"How many industries with a 6-digit NAICS code have a total recordable cases incidence rate above "
                f"{rnd(t, 1)}? Answer with only the integer.", str(c))

    def q_lookup():
        x = rng.choice([x for x in leaves if x["dj"] != x["away"] + x["tr"]])
        return (f"What is the combined incidence rate of cases with days away from work, job restriction, or transfer "
                f"in the {x['name']} industry, as reported in the table (not summed from its components)? Answer with "
                f"only the number using exactly two decimal places.", rnd(x["dj"], 2))

    def q_group_top():
        sname, scode = rng.choice(subs)
        x = top([g for g in groups if g["parent"] == scode], key=lambda g: g["total"])
        return (f"Among the 4-digit NAICS industry groups under '{sname}' (NAICS {scode}), which has the highest total "
                f"recordable cases incidence rate? Answer with only the industry name exactly as it appears in the "
                f"table.", x["name"])

    def q_code():
        x = top(leaves, key=lambda x: x["other"])
        return ("Which 6-digit NAICS code belongs to the industry with the highest 'Other recordable cases' incidence "
                "rate? Answer with only the NAICS code.", x["code"])

    def q_avg():
        g = rng.choice([g for g in groups if sum(x["parent"] == g["code"] for x in leaves) >= 3])
        v = [x["total"] for x in leaves if x["parent"] == g["code"]]
        return (f"What is the average total recordable cases incidence rate across the 6-digit industries in the "
                f"'{g['name']}' industry group (NAICS {g['code']})? Answer with only the number rounded to two "
                f"decimal places.", rnd(sum(v) / len(v), 2))

    return wb, [q_count, q_lookup, q_group_top, q_code, q_avg]


G_SECTORS = [("Food, beverages & tobacco", "10-12"), ("Textiles, wearing apparel & leather", "13-15"),
             ("Wood & paper products, & printing", "16-18"), ("Coke & refined petroleum", "19"),
             ("Chemicals, Pharmaceuticals", "20-21"), ("Rubber, plastics & non-metallic minerals", "22-23"),
             ("Basic metals & metal products", "24-25"), ("Computer etc products, Electrical equipment", "26-27"),
             ("Machinery & equipment n.e.c.", "28"), ("Transport equipment", "29-30"),
             ("Other manufacturing & repair", "31-33")]


def header_lines(name, k=4):
    words = name.split()
    n = len(words)
    if n <= k:
        return [" "] * (k - n) + [" " + w for w in words]
    sizes = [n // k + (i < n % k) for i in range(k)]
    out, i = [], 0
    for s in sizes:
        out.append(" " + " ".join(words[i:i + s]))
        i += s
    return out


def gen_quarterly(rng):
    chosen = set(rng.sample(G_SECTORS, rng.randint(6, 9)))
    secs = [s for s in G_SECTORS if s in chosen]
    y0, y1 = rng.randint(1975, 1995), rng.randint(2015, 2023)
    years = list(range(y0, y1 + 1))
    base = rng.randint(max(y0, 2010), y1)
    ann, qtr, level = {}, {}, {}
    for name, _ in secs:
        raw, v, g = {}, 1.0, rng.uniform(-.01, .04)
        for y in years:
            v *= math.exp(g + rng.gauss(0, .045))
            raw[y] = v
        for y in years:
            ann[name, y] = q(F(raw[y] / raw[base] * 100), 1)
            for k in range(1, 5):
                qtr[name, y, k] = q(F(float(ann[name, y]) * rng.uniform(.97, 1.03)), 1)
        level[name] = num(rng, 30, 180, 1)
    ids = rng.sample([f"DJ{a}{b}" for a in string.digits + string.ascii_uppercase[:6] for b in string.ascii_uppercase],
                     len(secs))

    wb, sh = workbook(f"Table {rng.randint(1, 4)}")
    r = sh.row([rng.randint(1, 4), " Output per job: Manufacturing subsections"])
    sh.row([None, " United Kingdom"])
    sh.merge(r, 1, r + 1, 1)
    sh.row([None])
    lines = [header_lines(n) for n, _ in secs]
    for i in range(4):
        sh.row([None, None] + [ln[i] for ln in lines])
    sh.row([" Divisions ", " "] + [f"  {c} " for _, c in secs])
    sh.row(["Level (GBP thousand)"])
    sh.row([None, base] + [level[n] for n, _ in secs])
    sh.row(["Indices", f"{base}=100"])
    sh.row([None, None] + ids)
    for y in years:
        sh.row([None, y] + [ann[n, y] for n, _ in secs])
    sh.row([None])
    sh.row(["Quarterly"])
    for y in years:
        for k in range(1, 5):
            sh.row([None, f"{y} Q{k}"] + [qtr[n, y, k] for n, _ in secs])
    sh.row([None])
    sh.row(["Source: synthetic labour productivity series. Output per job = gross value added / number of jobs."])

    names = [n for n, _ in secs]
    quarters = [(y, k) for y in years for k in range(1, 5)]

    def q_count():
        a, b = rng.sample(names, 2)
        c = sum(ann[a, y] > ann[b, y] for y in years)
        return (f"Using the annual rows of the index series (rows labelled with a year, not with a quarter), in how "
                f"many years was the output per job index of '{a}' strictly higher than that of '{b}'? Answer with "
                f"only the integer.", str(c))

    def q_max_year():
        a = rng.choice(names)
        y = top(years, key=lambda y: ann[a, y])
        return (f"Using the annual figures (rows labelled with a year, not with a quarter), in which year did the "
                f"'{a}' sector have the highest output per job index? Answer with only the year.", str(y))

    def q_low_quarter():
        a = rng.choice(names)
        y, k = top(quarters, key=lambda t: qtr[a, t[0], t[1]], low=True)
        return (f"Using the quarterly rows, in which quarter was the output per job index of '{a}' lowest? Answer in "
                f"the form `YYYY Qn`, for example `1999 Q3`.", f"{y} Q{k}")

    def q_lookup():
        a, (y, k) = rng.choice(names), rng.choice(quarters)
        return (f"What was the output per job index of '{a}' in {y} Q{k}? Answer with only the number using exactly "
                f"one decimal place.", rnd(qtr[a, y, k], 1))

    def q_level():
        a = rng.choice(names)
        return (f"What was the level of output per job, in GBP thousand, for '{a}' in {base}? Answer with only the "
                f"number using exactly one decimal place.", rnd(level[a], 1))

    def q_change():
        a = rng.choice(names)
        ya, yb = sorted(rng.sample(years, 2))
        return (f"Using the annual figures, by how many index points did output per job in '{a}' change from {ya} to "
                f"{yb}? Answer with only the number using exactly one decimal place, with a leading minus sign if it "
                f"fell.", rnd(ann[a, yb] - ann[a, ya], 1))

    return wb, [q_count, q_max_year, q_low_quarter, q_lookup, q_level, q_change]


FV_ITEMS = [("Municipal bonds", 2), ("Municipal auction rate securities", 3), ("Corporate notes", 2),
            ("U.S. agency securities", 2), ("Asset-backed securities", 2)]


def gen_segments(rng):
    nseg = rng.randint(2, 4)
    fy = rng.randint(2012, 2024)
    fys = [fy, fy - 1, fy - 2]
    seg = []
    for _ in range(nseg):
        p0, s0 = rng.randint(150, 2500), rng.randint(200, 2500)
        gp, gs, mg = rng.uniform(-.08, .12), rng.uniform(-.03, .15), rng.uniform(.25, .65)
        d = {}
        for k, f in enumerate(reversed(fys)):
            prod = round(p0 * (1 + gp) ** k * rng.uniform(.97, 1.03))
            serv = round(s0 * (1 + gs) ** k * rng.uniform(.97, 1.03))
            d[f] = dict(prod=prod, serv=serv, total=prod + serv, op=round((prod + serv) * mg * rng.uniform(.9, 1.1)))
        seg.append(d)
    grand = {f: sum(s[f]["total"] for s in seg) for f in fys}

    def pct(a, b):
        return rnd(F(a - b, b) * 100, 0, check=False) + "%"

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for i, s in enumerate(seg):
        sh = Sheet(wb.create_sheet(f"table_{i}"))
        sh.row(["(Dollars in millions)"] + [f"Fiscal{f}" for f in fys] + [f"{fy}-{fy - 1}% Change",
                                                                         f"{fy - 1}-{fy - 2}% Change"])
        sh.row(["Product revenue"] + [money(s[f]["prod"], True) for f in fys])
        sh.row(["Service and other revenue"] + [money(s[f]["serv"]) for f in fys])
        sh.row(["Total segment revenue"] + [money(s[f]["total"], True) for f in fys]
               + [pct(s[fy]["total"], s[fy - 1]["total"]), pct(s[fy - 1]["total"], s[fy - 2]["total"])])
        sh.row(["% of total revenue"] + [rnd(F(s[f]["total"] * 100, grand[f]), 0, False) + "%" for f in fys])
        sh.row(["Segment operating income"] + [money(s[f]["op"], True) for f in fys]
               + [pct(s[fy]["op"], s[fy - 1]["op"]), pct(s[fy - 1]["op"], s[fy - 2]["op"])])
        sh.row(["% of related revenue"] + [rnd(F(s[f]["op"] * 100, s[f]["total"]), 0, False) + "%" for f in fys])

    fv = None
    if rng.random() < .7:
        dates = [fy, fy - 1]
        fv = {}
        for d in dates:
            cash = rng.randint(300, 1500)
            items = {n: (rng.randint(0, 800) if rng.random() < .85 else 0, lv) for n, lv in FV_ITEMS}
            eq = rng.randint(0, 60) if rng.random() < .6 else 0
            lv = {1: cash + eq, 2: sum(v for v, l in items.values() if l == 2),
                  3: sum(v for v, l in items.values() if l == 3)}
            fv[d] = dict(cash=cash, items=items, eq=eq, lv=lv, total=sum(lv.values()),
                         notes=rng.randint(300, 900))
        sh = Sheet(wb.create_sheet(f"table_{nseg}"))
        r = sh.row([None, f"At July 31, {dates[0]}", None, None, None, f"At July 31, {dates[1]}"])
        sh.merge(r, 2, r, 5)
        sh.merge(r, 6, r, 9)
        sh.row(["(In millions)"] + ["Level 1", "Level 2", "Level 3", "TotalFair Value"] * 2)
        sh.row(["Assets:"])

        def line(label, per_date, dollar=False):
            sh.row([label] + sum(([money(v, dollar) for v in per_date(d)] for d in dates), []))

        line("Cash equivalents, primarily money market funds",
             lambda d: [fv[d]["cash"], 0, 0, fv[d]["cash"]], True)
        sh.row(["Available-for-sale debt securities:"])
        for n, l in FV_ITEMS:
            line(n, lambda d, n=n, l=l: [fv[d]["items"][n][0] if l == k else 0 for k in (1, 2, 3)]
                 + [fv[d]["items"][n][0]])
        line("Available-for-sale corporate equity securities", lambda d: [fv[d]["eq"], 0, 0, fv[d]["eq"]])
        line("Total available-for-sale securities",
             lambda d: [fv[d]["eq"], fv[d]["lv"][2], fv[d]["lv"][3], fv[d]["eq"] + fv[d]["lv"][2] + fv[d]["lv"][3]])
        line("Total assets measured at fair value on a recurring basis",
             lambda d: [fv[d]["lv"][1], fv[d]["lv"][2], fv[d]["lv"][3], fv[d]["total"]], True)
        sh.row(["Liabilities:"])
        line("Senior notes -1", lambda d: [0, fv[d]["notes"], 0, fv[d]["notes"]], True)
        sh.row(["(1) The fair value of senior notes is based on observable market prices in less active markets."])

    def q_total():
        f = rng.choice(fys)
        return (f"In Fiscal{f}, what was the combined Total segment revenue of all the segments reported in this "
                f"workbook, in millions of dollars? Answer with only the integer, without commas or units.",
                str(grand[f]))

    def q_which():
        f = rng.choice(fys)
        label, k = rng.choice([("Service and other revenue", "serv"), ("Product revenue", "prod"),
                               ("Segment operating income", "op")])
        i = top(range(nseg), key=lambda i: seg[i][f][k])
        return (f"Which sheet reports the {'higher' if nseg == 2 else 'highest'} {label} in Fiscal{f}, and what was "
                f"that value in millions of dollars? Answer as `sheet name, integer value`, separated by a comma and "
                f"one space, without commas inside the number and without units.", f"table_{i}, {seg[i][f][k]}")

    def q_fv_change():
        if not fv:
            raise Retry
        return (f"By how much did the total assets measured at fair value on a recurring basis change between July "
                f"31, {fy - 1} and July 31, {fy}, in millions of dollars? Answer with only the integer, with a minus "
                f"sign if it decreased, without commas or units.", str(fv[fy]["total"] - fv[fy - 1]["total"]))

    def q_fv_share():
        if not fv:
            raise Retry
        d = rng.choice([fy, fy - 1])
        return (f"What share of the total assets measured at fair value on a recurring basis at July 31, {d} was "
                f"classified as Level 2? Answer as a percentage with exactly one decimal place, without the % sign.",
                rnd(F(fv[d]["lv"][2] * 100, fv[d]["total"]), 1))

    def q_margin():
        i, f = rng.randrange(nseg), rng.choice(fys)
        return (f"For the segment reported in sheet table_{i}, what was the segment operating income as a percentage "
                f"of total segment revenue in Fiscal{f}, computed from the dollar amounts? Answer with only the number "
                f"using exactly one decimal place.", rnd(F(seg[i][f]["op"] * 100, seg[i][f]["total"]), 1))

    def q_growth():
        i = rng.randrange(nseg)
        label, k = rng.choice([("Product revenue", "prod"), ("Service and other revenue", "serv")])
        a, b = seg[i][fy - 2][k], seg[i][fy][k]
        return (f"By what percentage did {label} in sheet table_{i} change from Fiscal{fy - 2} to Fiscal{fy}? Answer "
                f"with only the number using exactly one decimal place, with a leading minus sign if it decreased.",
                rnd(F(b - a, a) * 100, 1))

    return wb, [q_total, q_which, q_fv_change, q_fv_share, q_margin, q_growth]


RATINGS = ["AAA/Aaa to AA-/Aa3", "A+/A1 to A-/A3", "BBB+/Baa1 to BBB-/Baa3", "BB+/Ba1 to B-/B3", "CCC+/Caa1 and below"]
FUNDING = ["Wholesale funding", "Client cash management", "Obligations of Firm-administered multi-seller conduits(a)",
           "Other borrowed funds", "Securities sold under agreements to repurchase", "Securities loaned(b)",
           "Senior notes", "Trust preferred securities", "Subordinated debt", "Structured notes"]


def gen_bank(rng):
    Y = rng.randint(2012, 2024)
    c = dict(tse=rng.randint(150000, 300000), ps=rng.randint(5000, 30000), gw=rng.randint(20000, 50000),
             oia=rng.randint(300, 3000), dtl=rng.randint(1000, 4000), adj=rng.randint(500, 3000),
             ot1=rng.randint(100, 1000), ltd=rng.randint(10000, 20000), allow=rng.randint(3000, 6000),
             oth=rng.randint(50, 500))
    c["ce"] = c["tse"] - c["ps"]
    c["cet1"] = c["ce"] - c["gw"] - c["oia"] + c["dtl"] - c["adj"]
    c["t1"] = c["cet1"] + c["ps"] - c["ot1"]
    c["t2"] = c["ltd"] + c["allow"] + c["oth"]
    c["tc"] = c["t1"] + c["t2"]
    exp = {(b, y): rng.randint(500, 15000) for b in RATINGS for y in (Y, Y - 1)}
    etot = {y: sum(exp[b, y] for b in RATINGS) for y in (Y, Y - 1)}
    fund = {(f, k): (0 if f == "Client cash management" and rng.random() < .5 else rng.randint(1000, 160000))
            for f in FUNDING for k in ("end", "end_p", "avg", "avg_p")}

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sh = Sheet(wb.create_sheet("table_0"))
    sh.row(["(in millions)", f"December 31, {Y}"])
    for label, v, dollar in [("Total stockholders’ equity", c["tse"], True), ("Less: Preferred stock", c["ps"], False),
                             ("Common stockholders’ equity", c["ce"], False), ("Less:", None, 0),
                             ("Goodwill", c["gw"], False), ("Other intangible assets", c["oia"], False),
                             ("Add:", None, 0), ("Deferred tax liabilities(a)", c["dtl"], False),
                             ("Less: Other CET1 capital adjustments", c["adj"], False),
                             ("Standardized/Advanced CET1 capital", c["cet1"], True), ("Preferred stock", c["ps"], False),
                             ("Less:", None, 0), ("Other Tier 1 adjustments(b)", c["ot1"], False),
                             ("Standardized/Advanced Tier 1 capital", c["t1"], True),
                             ("Long-term debt and other instruments qualifying as Tier 2 capital", c["ltd"], True),
                             ("Qualifying allowance for credit losses", c["allow"], False), ("Other", c["oth"], False),
                             ("Standardized Tier 2 capital", c["t2"], True),
                             ("Standardized total capital", c["tc"], True)]:
        sh.row([label, None if v is None else money(v, dollar)])
    sh.row(["(a) Represents deferred tax liabilities related to tax-deductible goodwill and intangibles."])
    sh.row(["(b) Includes the deduction associated with the permissible holdings of covered funds."])

    sh = Sheet(wb.create_sheet("table_1"))
    r = sh.row(["Rating equivalent", str(Y), None, f"{Y - 1}(a)", None])
    sh.merge(r, 2, r, 3)
    sh.merge(r, 4, r, 5)
    sh.row(["December 31,(in millions, except ratios)", "Exposure net of all collateral",
            "% of exposure netof all collateral", "Exposure net of all collateral", "% of exposure netof all collateral"])
    for i, b in enumerate(RATINGS):
        cells = []
        for y in (Y, Y - 1):
            p = rnd(F(exp[b, y] * 100, etot[y]), 0, False)
            cells += [money(exp[b, y], i == 0), p + "%" if i == 0 else p]
        sh.row([b] + cells)
    sh.row(["Total", money(etot[Y], True), "100%", money(etot[Y - 1], True), "100%"])

    sh = Sheet(wb.create_sheet("table_2"))
    r = sh.row(["As of or for the year ended December 31,", None, None, "Average", None])
    sh.merge(r, 4, r, 5)
    sh.row(["(in millions)", str(Y), str(Y - 1), str(Y), str(Y - 1)])
    for f in FUNDING:
        if f == "Wholesale funding":
            sh.row(["Commercial paper:"])
        if f == "Securities sold under agreements to repurchase":
            sh.row(["Securities loaned or sold under agreements to repurchase:"])
        sh.row([f] + [money(fund[f, k], True) for k in ("end", "end_p", "avg", "avg_p")])
    sh.row(["(a) Includes conduits consolidated by the Firm.", None])
    sh.row(["(b) Prior-period amounts have been revised to conform with the current presentation.", None])

    clean = [f for f in FUNDING if "(" not in f and f != "Client cash management"]
    cap = {"Preferred stock": c["ps"], "Goodwill": c["gw"], "Other intangible assets": c["oia"],
           "Qualifying allowance for credit losses": c["allow"]}

    def q_ratio():
        f, z = rng.choice(clean), rng.choice(list(cap))
        return (f"What is the ratio of {f} at December 31, {Y} to {z} at December 31, {Y}? Answer with only the "
                f"number rounded to three decimal places.", rnd(F(fund[f, "end"], cap[z]), 3))

    def q_share():
        b, y = rng.choice(RATINGS), rng.choice([Y, Y - 1])
        return (f"What percentage of the total exposure net of all collateral at December 31, {y} came from the "
                f"'{b}' rating equivalent, computed from the dollar amounts? Answer with only the number using exactly "
                f"one decimal place.", rnd(F(exp[b, y] * 100, etot[y]), 1))

    def q_avg_end():
        f = rng.choice(clean)
        return (f"For {f}, how much larger was the {Y} average balance than the balance at December 31, {Y}, in "
                f"millions of dollars? Answer with only the integer, with a minus sign if the average was smaller, "
                f"without commas or units.", str(fund[f, "avg"] - fund[f, "end"]))

    def q_sum():
        return (f"What is the sum of Goodwill and Other intangible assets at December 31, {Y}, in millions of dollars? "
                f"Answer with only the integer, without commas or units.", str(c["gw"] + c["oia"]))

    def q_drop():
        dec = {b: exp[b, Y - 1] - exp[b, Y] for b in RATINGS}
        b = top(RATINGS, key=dec.get)
        if dec[b] <= 0:
            raise Retry
        return (f"Which rating equivalent category had the largest decrease in exposure net of all collateral from "
                f"{Y - 1} to {Y}? Answer with only the category exactly as it appears in the table.", b)

    def q_capital_gap():
        return (f"What is the difference between Standardized total capital and Standardized/Advanced CET1 capital "
                f"at December 31, {Y}, in millions of dollars? Answer with only the integer, without commas or units.",
                str(c["tc"] - c["cet1"]))

    return wb, [q_ratio, q_share, q_avg_end, q_sum, q_drop, q_capital_gap]


FIRST = ["An", "Binh", "Chi", "Dung", "Giang", "Ha", "Hai", "Hieu", "Hoa", "Hung", "Khanh", "Lan", "Linh", "Long",
         "Mai", "Minh", "Nam", "Nga", "Phuc", "Quan", "Son", "Thao", "Trang", "Tuan", "Vy", "Yen", "Duc", "Huong",
         "Tam", "Thu"]
MIDDLE = ["Van", "Thi", "Minh", "Ngoc", "Gia", "Hoang", "Bao", "Anh", "Quoc", "Thanh", "Duc", "Kim"]
LAST = ["Nguyen", "Tran", "Le", "Pham", "Hoang", "Vu", "Vo", "Dang", "Bui", "Do", "Ho", "Ngo", "Duong", "Ly", "Phan",
        "Truong"]
ROSTER = [dict(sheet="People", event="Admission", ev="admission", score="Score", sc="score", verb="admitted"),
          dict(sheet="Staff", event="Hire", ev="hire", score="Rating", sc="rating", verb="hired"),
          dict(sheet="Members", event="Joining", ev="joining", score="Points", sc="points", verb="joined")]
ORD = {1: "", 2: "second-", 3: "third-"}
WORDS = {5: "five", 6: "six", 7: "seven", 8: "eight", 9: "nine", 10: "ten"}
DEPTS = ["Sales", "IT", "HR", "Finance", "Marketing"]


def gen_roster(rng):
    v = rng.choice(ROSTER)
    with_dept = rng.random() < .5
    people = []
    for i in range(rng.randint(15, 40)):
        ev = dt.datetime(2015, 1, 1) + dt.timedelta(days=rng.randint(0, 11 * 365))
        birth = ev - dt.timedelta(days=rng.randint(18 * 365, 45 * 365))
        people.append(dict(no=i + 1, first=rng.choice(FIRST), middle=rng.choice(MIDDLE), last=rng.choice(LAST),
                           birth=birth, ev=ev, score=rng.randint(50, 100), dept=rng.choice(DEPTS)))

    wb, sh = workbook(v["sheet"])
    r = sh.row(["No", "Name", None, None, "Date of", None, v["score"]] + (["Department"] if with_dept else []))
    sh.row([None, "First", "Middle", "Last", "Birth", v["event"], None])
    for a, b, c, d in [(r, 2, r, 4), (r, 5, r, 6), (r, 7, r + 1, 7), (r, 1, r + 1, 1)] + (
            [(r, 8, r + 1, 8)] if with_dept else []):
        sh.merge(a, b, c, d)
    for p in people:
        r = sh.row([p["no"], p["first"], p["middle"], p["last"], p["birth"], p["ev"], p["score"]]
                   + ([p["dept"]] if with_dept else []))
        for col in (5, 6):
            sh.ws.cell(r, col).number_format = "yyyy-mm-dd"

    def nth(lst, key, k):
        s = sorted(lst, key=key)
        if len({key(x) for x in s[:k + 1]}) < k + 1:
            raise Retry
        return s[k - 1]

    def q_event_age():
        k, young = rng.choice([1, 2, 3]), rng.random() < .6
        p = nth(people, lambda p: (p["ev"] - p["birth"]).days * (1 if young else -1), k)
        word = ORD[k] + ("youngest" if young else "oldest")
        return (f"Who was the {word} person on their {v['ev']} date? Answer with only the last name.", p["last"])

    def q_score_by_age():
        k, old = rng.choice([1, 2, 3]), rng.random() < .6
        p = nth(people, lambda p: p["birth"].toordinal() * (1 if old else -1), k)
        word = ORD[k] + ("oldest" if old else "youngest")
        return (f"What {v['sc']} did the {word} person receive? Answer with only the integer {v['sc']}.",
                str(p["score"]))

    def q_best_among_young():
        n = rng.randint(5, 10)
        s = sorted(people, key=lambda p: p["birth"], reverse=True)
        if s[n - 1]["birth"] == s[n]["birth"]:
            raise Retry
        return (f"Among the {WORDS[n]} youngest people, what was the highest {v['sc']}? Answer with only the integer "
                f"{v['sc']}.", str(max(p["score"] for p in s[:n])))

    def q_count_year():
        y = rng.choice(people)["ev"].year
        return (f"How many people were {v['verb']} in {y}? Answer with only the integer.",
                str(sum(p["ev"].year == y for p in people)))

    def q_avg_born():
        y = rng.randint(1975, 2000)
        sel = [p["score"] for p in people if p["birth"].year < y]
        if len(sel) < 2:
            raise Retry
        return (f"What is the average {v['sc']} of the people born before {y}? Answer with only the number rounded to "
                f"two decimal places.", rnd(F(sum(sel), len(sel)), 2))

    def q_full_name():
        p = top(people, key=lambda p: p["score"])
        return (f"Who has the highest {v['sc']}? Answer with the full name in the form `First Middle Last`.",
                f"{p['first']} {p['middle']} {p['last']}")

    def q_dept():
        if not with_dept:
            raise Retry
        d = rng.choice(DEPTS)
        sel = [p["score"] for p in people if p["dept"] == d]
        if not sel:
            raise Retry
        return (f"What is the average {v['sc']} of the people in the {d} department? Answer with only the number "
                f"rounded to two decimal places.", rnd(F(sum(sel), len(sel)), 2))

    return wb, [q_event_age, q_score_by_age, q_best_among_young, q_count_year, q_avg_born, q_full_name, q_dept]


REGIONS = [{"vi": "Canada", "fr": "Canada", "ko": "캐나다", "zh": "加拿大"},
           {"vi": "Ontario", "fr": "Ontario", "ko": "온타리오", "zh": "安大略省"},
           {"vi": "Quebec", "fr": "Québec", "ko": "퀘벡", "zh": "魁北克省"},
           {"vi": "Alberta", "fr": "Alberta", "ko": "앨버타", "zh": "艾伯塔省"},
           {"vi": "British Columbia", "fr": "Colombie-Britannique", "ko": "브리티시컬럼비아", "zh": "不列颠哥伦比亚省"}]
SMALL = {
    "care": {
        "vi": dict(title="dân số từ 12 tuổi trở lên đã nhận hoặc cần được chăm sóc, {region}, {year}",
                   rowhdr="loại nhu cầu", colgrp="tình trạng đáp ứng",
                   cols=["chưa được đáp ứng", "đáp ứng một phần", "đáp ứng đầy đủ"], unit="%",
                   rows=["bất kỳ", "thông tin về vấn đề, cách điều trị hoặc dịch vụ hiện có", "thuốc men",
                         "tư vấn, trị liệu hoặc hỗ trợ về các vấn đề trong quan hệ cá nhân",
                         "hỗ trợ chăm sóc cá nhân", "hỗ trợ việc nhà"],
                   qa="Trong các loại nhu cầu chăm sóc (không tính dòng `bất kỳ` và bỏ qua dòng có giá trị `..`), "
                      "tổng tỷ lệ chưa được đáp ứng và đáp ứng một phần cao nhất là bao nhiêu phần trăm? Chỉ trả lời "
                      "bằng số với đúng một chữ số thập phân.",
                   qb="Loại nhu cầu nào (không tính dòng `bất kỳ`) có tỷ lệ được đáp ứng đầy đủ cao nhất? Chỉ trả lời "
                      "bằng tên loại nhu cầu đúng như trong bảng."),
        "fr": dict(title="population âgée de 12 ans et plus ayant reçu ou eu besoin de soins, {region}, {year}",
                   rowhdr="type de besoin", colgrp="état du besoin",
                   cols=["non satisfait", "partiellement satisfait", "entièrement satisfait"], unit="%",
                   rows=["tout besoin", "information sur les problèmes, les traitements ou les services offerts",
                         "médicaments", "counseling, thérapie ou aide pour des problèmes de relations personnelles",
                         "aide pour les soins personnels", "aide pour les tâches ménagères"],
                   qa="Parmi les types de besoins (sans compter la ligne `tout besoin` et en ignorant les lignes "
                      "contenant `..`), quelle est la somme la plus élevée des pourcentages de besoins non satisfaits "
                      "et partiellement satisfaits ? Répondez uniquement par le nombre avec une décimale, sans le "
                      "symbole %.",
                   qb="Quel type de besoin (hors ligne `tout besoin`) présente le pourcentage de besoins entièrement "
                      "satisfaits le plus élevé ? Répondez uniquement par le nom du type de besoin tel qu'il figure "
                      "dans le tableau."),
        "ko": dict(title="{year}년 {region}, 치료·돌봄을 받았거나 필요로 한 12세 이상 인구",
                   rowhdr="필요 유형", colgrp="충족 상태", cols=["미충족", "부분 충족", "완전 충족"], unit="%",
                   rows=["전체", "문제, 치료 또는 이용 가능한 서비스에 관한 정보", "약물",
                         "상담, 치료 또는 대인 관계 문제에 대한 도움", "개인 돌봄 지원", "가사 지원"],
                   qa="필요 유형 중(`전체` 행 제외, `..` 값이 있는 행은 무시) 미충족 비율과 부분 충족 비율의 합이 가장 "
                      "높은 값은 몇 퍼센트입니까? 소수점 첫째 자리까지, % 기호 없이 숫자만 답하십시오.",
                   qb="`전체` 행을 제외하고 완전 충족 비율이 가장 높은 필요 유형은 무엇입니까? 표에 표시된 그대로 필요 "
                      "유형 이름만 답하십시오."),
        "zh": dict(title="{year}年{region}接受或需要护理服务的12岁及以上人口",
                   rowhdr="需求类型", colgrp="需求满足状况", cols=["未满足", "部分满足", "完全满足"], unit="%",
                   rows=["任何需求", "有关问题、治疗或现有服务的信息", "药物", "心理咨询、治疗或人际关系问题帮助",
                         "个人护理帮助", "家务帮助"],
                   qa="在各需求类型中（不包括“任何需求”行，并忽略含有“..”的行），未满足与部分满足比例之和的最大值是"
                      "多少？只回答数字，保留一位小数，不带%符号。",
                   qb="除“任何需求”行外，哪种需求类型的完全满足比例最高？只回答表中所示的需求类型名称。"),
    },
    "income": {
        "vi": dict(title="thu nhập tiền lương và tiền công trung vị của {grp} có việc làm, độ tuổi 25 đến 54, theo "
                         "trình độ học vấn cao nhất và giới tính, {region}, {year}",
                   grps=["người Métis", "người thuộc Các Dân tộc Thứ nhất", "người Inuit"],
                   cols=["nam", "nữ"], unit="đô la",
                   rows=["tổng thu nhập tiền lương và tiền công trung vị", "chưa tốt nghiệp trung học phổ thông",
                         "bằng tốt nghiệp trung học phổ thông hoặc tương đương", "có học một phần sau trung học",
                         "chứng chỉ hoặc văn bằng sau trung học dưới bậc cử nhân",
                         "chứng chỉ, văn bằng hoặc bằng cấp đại học"],
                   qa="Giữa các nhóm trình độ học vấn (không tính dòng tổng), chênh lệch thu nhập tiền lương và tiền "
                      "công trung vị giữa nam và nữ lớn nhất là bao nhiêu đô la? Chỉ trả lời bằng số nguyên, không có "
                      "dấu phẩy.",
                   qb="Ở nhóm `{row}`, thu nhập tiền lương và tiền công trung vị của nữ bằng bao nhiêu phần trăm so "
                      "với nam? Làm tròn đến một chữ số thập phân, chỉ trả lời bằng số, không kèm ký hiệu %."),
        "fr": dict(title="revenu d'emploi médian {grp} occupés, âgés de 25 à 54 ans, selon le plus haut niveau de "
                         "scolarité et le sexe, {region}, {year}",
                   grps=["des Métis", "des Premières Nations", "des Inuits"], cols=["hommes", "femmes"],
                   unit="dollars",
                   rows=["revenu d'emploi médian, total", "aucun diplôme d'études secondaires",
                         "diplôme d'études secondaires ou l'équivalent", "études postsecondaires partielles",
                         "certificat ou diplôme d'études postsecondaires inférieur au baccalauréat",
                         "certificat, diplôme ou grade universitaire"],
                   qa="Parmi les niveaux de scolarité (sans compter la ligne du total), quel est le plus grand écart, "
                      "en dollars, entre le revenu d'emploi médian des hommes et celui des femmes ? Répondez uniquement "
                      "par un nombre entier, sans espace ni virgule.",
                   qb="Pour le niveau `{row}`, le revenu d'emploi médian des femmes représente quel pourcentage de "
                      "celui des hommes ? Arrondissez à une décimale et répondez uniquement par le nombre, sans le "
                      "symbole %."),
        "ko": dict(title="{year}년 {region}, 최종 학력 및 성별 25~54세 {grp} 취업자의 중위 임금·급여 소득",
                   grps=["메티스", "퍼스트 네이션", "이누이트"], cols=["남성", "여성"], unit="달러",
                   rows=["중위 임금·급여 소득 전체", "고등학교 미졸업", "고등학교 졸업 또는 동등 학력",
                         "중등 후 교육 일부 이수", "학사 미만 중등 후 교육 수료증 또는 졸업장",
                         "대학 수료증, 졸업장 또는 학위"],
                   qa="학력 집단 중(합계 행 제외) 남성과 여성의 중위 임금·급여 소득 차이가 가장 큰 값은 몇 달러입니까? "
                      "쉼표 없이 정수로만 답하십시오.",
                   qb="`{row}` 집단에서 여성의 중위 임금·급여 소득은 남성의 몇 퍼센트입니까? 소수점 첫째 자리까지 "
                      "반올림하여 % 기호 없이 숫자만 답하십시오."),
        "zh": dict(title="{year}年{region}25至54岁在业{grp}按最高学历和性别划分的工资薪金收入中位数",
                   grps=["梅蒂斯人", "第一民族", "因纽特人"], cols=["男性", "女性"], unit="加元",
                   rows=["工资薪金收入中位数总计", "未完成高中", "高中毕业或同等学历", "接受过部分高等教育",
                         "低于学士学位的高等教育证书或文凭", "大学证书、文凭或学位"],
                   qa="在各学历组别中（不含总计行），男性与女性工资薪金收入中位数的最大差距是多少加元？只回答整数，"
                      "不加逗号。",
                   qb="在“{row}”组别中，女性的工资薪金收入中位数相当于男性的百分之多少？保留一位小数，只回答数字，"
                      "不带%符号。"),
    },
    "employment": {
        "vi": dict(title="tỷ lệ có việc làm của dân số từ 25 đến 64 tuổi theo tình trạng hôn nhân và bản sắc bản "
                         "địa, {region}, {year}",
                   rowhdr="tình trạng hôn nhân", colgrp="bản sắc bản địa",
                   cols=["tổng", "người bản địa", "người không phải bản địa"], unit="%",
                   rows=["tổng, mọi tình trạng hôn nhân", "đã kết hôn", "sống chung không kết hôn", "ly thân",
                         "ly hôn", "góa", "độc thân (chưa từng kết hôn hợp pháp)"],
                   qa="Đối với nhóm `{row}`, tỷ lệ có việc làm của người bản địa thấp hơn người không phải bản địa "
                      "bao nhiêu điểm phần trăm? Chỉ trả lời bằng số với đúng một chữ số thập phân.",
                   qb="Trong các tình trạng hôn nhân (không tính dòng tổng), nhóm nào có chênh lệch tỷ lệ có việc "
                      "làm giữa người không phải bản địa và người bản địa lớn nhất? Chỉ trả lời bằng tên nhóm đúng "
                      "như trong bảng."),
        "fr": dict(title="taux d'emploi de la population âgée de 25 à 64 ans selon l'état matrimonial et "
                         "l'identité autochtone, {region}, {year}",
                   rowhdr="état matrimonial", colgrp="identité autochtone",
                   cols=["total", "Autochtones", "non-Autochtones"], unit="pourcentage",
                   rows=["total, tous les états matrimoniaux", "marié", "en union libre", "séparé", "divorcé", "veuf",
                         "célibataire (jamais légalement marié)"],
                   qa="Pour le groupe `{row}`, de combien de points de pourcentage le taux d'emploi des Autochtones "
                      "est-il inférieur à celui des non-Autochtones ? Répondez uniquement par le nombre avec une "
                      "décimale, sans le symbole %.",
                   qb="Parmi les états matrimoniaux (hors ligne du total), lequel présente le plus grand écart de "
                      "taux d'emploi entre les non-Autochtones et les Autochtones ? Répondez uniquement par le nom tel "
                      "qu'il figure dans le tableau."),
        "ko": dict(title="{year}년 {region}, 혼인 상태 및 원주민 정체성별 25~64세 인구의 고용률",
                   rowhdr="혼인 상태", colgrp="원주민 정체성", cols=["전체", "원주민", "비원주민"], unit="%",
                   rows=["전체, 모든 혼인 상태", "기혼", "사실혼", "별거", "이혼", "사별", "미혼(법적 혼인 경험 없음)"],
                   qa="`{row}` 집단에서 원주민의 고용률은 비원주민보다 몇 퍼센트포인트 낮습니까? 소수점 첫째 자리까지 "
                      "숫자만 답하십시오.",
                   qb="혼인 상태 중(전체 행 제외) 비원주민과 원주민의 고용률 차이가 가장 큰 집단은 무엇입니까? 표에 "
                      "표시된 그대로 이름만 답하십시오."),
        "zh": dict(title="{year}年{region}25至64岁人口按婚姻状况和原住民身份划分的就业率",
                   rowhdr="婚姻状况", colgrp="原住民身份", cols=["总计", "原住民", "非原住民"], unit="百分比",
                   rows=["总计，所有婚姻状况", "已婚", "同居", "分居", "离婚", "丧偶", "单身（从未合法结婚）"],
                   qa="对于“{row}”人群，原住民的就业率比非原住民低多少个百分点？只回答数字，保留一位小数。",
                   qb="在各婚姻状况中（不含总计行），哪一组非原住民与原住民的就业率差距最大？只回答表中所示的名称。"),
    },
    "court": {
        "vi": dict(title="số vụ án đã hoàn tất tại tòa án hình sự dành cho người trưởng thành, theo loại tội phạm, "
                         "{region}, {y0} đến {y1}",
                   rowhdr="năm", colgrp="loại tội phạm",
                   cols=["tổng số vụ án", "tội xâm phạm thân thể", "tội xâm phạm tài sản",
                         "tội chống lại hoạt động tư pháp", "các tội khác theo Bộ luật Hình sự"], unit="số vụ",
                   qa="So với năm {y0}, tổng số vụ án đã hoàn tất tại tòa án hình sự dành cho người trưởng thành năm "
                      "{y1} giảm bao nhiêu phần trăm? Làm tròn đến hai chữ số thập phân, chỉ trả lời bằng số, không "
                      "kèm ký hiệu %.",
                   qb="Trong năm nào số vụ án thuộc nhóm `{col}` cao nhất? Chỉ trả lời bằng năm đúng như trong bảng "
                      "(ví dụ 2005/2006)."),
        "fr": dict(title="causes réglées par les tribunaux de juridiction criminelle pour adultes, selon le type "
                         "d'infraction, {region}, {y0} à {y1}",
                   rowhdr="année", colgrp="type d'infraction",
                   cols=["total des causes", "crimes contre la personne", "crimes contre les biens",
                         "infractions contre l'administration de la justice", "autres infractions au Code criminel"],
                   unit="nombre",
                   qa="De quel pourcentage le nombre total de causes réglées par les tribunaux de juridiction "
                      "criminelle pour adultes a-t-il diminué entre {y0} et {y1} ? Répondez uniquement par le nombre "
                      "avec deux décimales, sans le symbole %.",
                   qb="Au cours de quelle année le nombre de causes de la catégorie `{col}` a-t-il été le plus élevé ? "
                      "Répondez uniquement par l'année telle qu'elle figure dans le tableau (par exemple 2005/2006)."),
        "ko": dict(title="{region} 성인 형사 법원에서 완료된 사건 수, 범죄 유형별, {y0}~{y1}",
                   rowhdr="연도", colgrp="범죄 유형",
                   cols=["전체 사건", "대인 범죄", "재산 범죄", "사법 운영 방해 범죄", "기타 형법 위반"], unit="건수",
                   qa="{y0} 대비 {y1}에 성인 형사 법원에서 완료된 전체 사건 수는 몇 퍼센트 감소했습니까? 소수점 둘째 "
                      "자리까지, % 기호 없이 숫자만 답하십시오.",
                   qb="`{col}` 사건 수가 가장 많았던 연도는 언제입니까? 표에 표시된 그대로 연도만 답하십시오"
                      "(예: 2005/2006)."),
        "zh": dict(title="{region}成人刑事法院已结案件数，按犯罪类型划分，{y0}至{y1}",
                   rowhdr="年度", colgrp="犯罪类型",
                   cols=["案件总数", "侵犯人身罪", "侵犯财产罪", "妨碍司法罪", "其他刑法罪行"], unit="件",
                   qa="与{y0}年度相比，{y1}年度成人刑事法院已结案件总数下降了百分之多少？保留两位小数，只回答数字，"
                      "不带%符号。",
                   qb="“{col}”案件数量最多的是哪一年度？只回答表中所示的年度（例如2005/2006）。"),
    },
}


def gen_small(rng, theme, lang):
    L = SMALL[theme][lang]
    region = rng.choice(REGIONS)[lang]
    year = rng.randint(2010, 2022)
    wb, sh = workbook("table")
    s1 = lambda x: rnd(x, 1, check=False)
    pool = []

    def header(title, rowhdr, colgrp, cols, unit):
        nc = 1 + len(cols)
        r = sh.row([title])
        sh.merge(r, 1, r, nc)
        r = sh.row([rowhdr, colgrp])
        sh.merge(r, 1, r + 2, 1)
        sh.merge(r, 2, r, nc)
        sh.row([None] + cols)
        r = sh.row([None, unit])
        sh.merge(r, 2, r, nc)

    if theme == "care":
        header(L["title"].format(region=region, year=year), L["rowhdr"], L["colgrp"], L["cols"], L["unit"])
        vals, missing = {}, None
        for i, name in enumerate(L["rows"]):
            u, p = num(rng, 5, 40, 1), num(rng, 3, 25, 1)
            vals[name] = (u, p, 100 - u - p)
        if rng.random() < .5:
            missing = rng.choice(L["rows"][1:])
        for name in L["rows"]:
            u, p, f = vals[name]
            sh.row([name, ".." if name == missing else s1(u), s1(p), s1(f)])
        rows = L["rows"][1:]

        def qa():
            cand = [r for r in rows if r != missing]
            r = top(cand, key=lambda r: vals[r][0] + vals[r][1])
            return L["qa"], rnd(vals[r][0] + vals[r][1], 1)

        def qb():
            return L["qb"], top(rows, key=lambda r: vals[r][2])
        pool = [qa, qb]

    elif theme == "income":
        grp = rng.choice(L["grps"])
        nc = 3
        r = sh.row([L["title"].format(grp=grp, region=region, year=year)])
        sh.merge(r, 1, r, nc)
        r = sh.row([None] + L["cols"])
        sh.merge(r, 1, r + 1, 1)
        r = sh.row([None, L["unit"]])
        sh.merge(r, 2, r, nc)
        vals = {}
        for i, name in enumerate(L["rows"]):
            m = rng.randint(30000 + 7000 * i, 45000 + 9000 * i) if i else rng.randint(45000, 65000)
            vals[name] = (m, round(m * rng.uniform(.55, .92)))
            sh.row([name, f"{vals[name][0]:,}", f"{vals[name][1]:,}"])
        rows = L["rows"][1:]

        def qa():
            r = top(rows, key=lambda r: vals[r][0] - vals[r][1])
            return L["qa"], str(vals[r][0] - vals[r][1])

        def qb():
            r = rng.choice(rows)
            return L["qb"].format(row=r), rnd(F(vals[r][1] * 100, vals[r][0]), 1)
        pool = [qa, qb]

    elif theme == "employment":
        header(L["title"].format(region=region, year=year), L["rowhdr"], L["colgrp"], L["cols"], L["unit"])
        vals = {}
        for name in L["rows"]:
            non = num(rng, 58, 88, 1)
            ind = non - num(rng, 2, 22, 1)
            vals[name] = (non - num(rng, .3, 2, 1), ind, non)
            sh.row([name] + [s1(x) for x in vals[name]])
        rows = L["rows"][1:]

        def qa():
            r = rng.choice(L["rows"])
            return L["qa"].format(row=r), rnd(vals[r][2] - vals[r][1], 1)

        def qb():
            return L["qb"], top(rows, key=lambda r: vals[r][2] - vals[r][1])
        pool = [qa, qb]

    else:
        s = rng.randint(2003, 2012)
        years = [f"{y}/{y + 1}" for y in range(s, s + 10)]
        y0, y1 = years[0], years[-1]
        header(L["title"].format(region=region, y0=y0, y1=y1), L["rowhdr"], L["colgrp"], L["cols"], L["unit"])
        base = [rng.randint(20000, 120000) for _ in range(4)]
        trend = [rng.uniform(-.035, .01) for _ in range(4)]
        vals = {}
        for k, y in enumerate(years):
            parts = [round(b * (1 + t) ** k * rng.uniform(.96, 1.04)) for b, t in zip(base, trend)]
            vals[y] = [sum(parts)] + parts
            sh.row([y] + [f"{v:,}" for v in vals[y]])
        if vals[y1][0] >= vals[y0][0]:
            raise Retry

        def qa():
            return L["qa"].format(y0=y0, y1=y1), rnd(F(vals[y0][0] - vals[y1][0], vals[y0][0]) * 100, 2)

        def qb():
            j = rng.randrange(5)
            return L["qb"].format(col=L["cols"][j]), top(years, key=lambda y: vals[y][j])
        pool = [qa, qb]
    return wb, pool


GEN = {"activity": gen_activity, "cpi": gen_cpi, "annual": gen_annual, "ethnic": gen_ethnic,
       "occupation": gen_occupation, "quartile": gen_quartile, "naics": gen_naics, "quarterly": gen_quarterly,
       "segments": gen_segments, "bank": gen_bank, "roster": gen_roster, "small": gen_small}
PLAN = {"activity": 7, "cpi": 7, "annual": 7, "ethnic": 4, "occupation": 3, "quartile": 6, "naics": 5,
        "quarterly": 5, "segments": 5, "bank": 4, "roster": 6, "small": 16}


def allocate(n):
    tot = sum(PLAN.values())
    raw = {k: v * n / tot for k, v in PLAN.items()}
    out = {k: int(x) for k, x in raw.items()}
    for k in sorted(raw, key=lambda k: raw[k] - out[k], reverse=True)[:n - sum(out.values())]:
        out[k] += 1
    return out


def build(kind, rng, extra, stats):
    for _ in range(100):
        try:
            wb, pool = GEN[kind](rng, **extra)
        except Retry:
            continue
        qs = []
        for fn in pool:
            for _ in range(8):
                try:
                    qs.append((fn.__name__, *fn()))
                    stats[kind, fn.__name__, "ok"] += 1
                    break
                except (Retry, IndexError, ValueError):
                    stats[kind, fn.__name__, "fail"] += 1
        if len(qs) >= 2:
            return wb, qs
    raise RuntimeError(f"không sinh được bảng {kind}")


def main():
    ap = argparse.ArgumentParser(description="Sinh bộ dữ liệu Excel QA tổng hợp")
    ap.add_argument("--out", default="data_x5")
    ap.add_argument("--files", type=int, default=75)
    ap.add_argument("--queries", type=int, default=150)
    ap.add_argument("--prefix", default="p5")
    ap.add_argument("--seed", type=int, default=2026)
    args = ap.parse_args()
    rng = random.Random(args.seed)

    jobs = []
    combos = [(t, lang) for t in SMALL for lang in ("vi", "fr", "ko", "zh")]
    rng.shuffle(combos)
    for kind, n in allocate(args.files).items():
        for i in range(n):
            extra = dict(zip(("theme", "lang"), combos[i % len(combos)])) if kind == "small" else {}
            jobs.append((kind, extra))
    rng.shuffle(jobs)

    out = Path(args.out)
    (out / "xlsx_files").mkdir(parents=True, exist_ok=True)
    stats, built = Counter(), []
    for i, (kind, extra) in enumerate(jobs, 1):
        wb, qs = build(kind, rng, extra, stats)
        name = f"{args.prefix}_t{i:02d}.xlsx"
        wb.save(out / "xlsx_files" / name)
        built.append([name, kind, qs, 2])

    def can_add(b):
        return b[3] < min(3, len(b[2])) if len(b[2]) >= 3 else False
    for _ in range(len(built) // 5):
        a, b = rng.sample(built, 2)
        if can_add(a) and b[3] > 1:
            a[3] += 1
            b[3] -= 1
    total = sum(b[3] for b in built)
    while total < args.queries:
        b = rng.choice([b for b in built if b[3] < len(b[2])])
        b[3] += 1
        total += 1
    while total > args.queries:
        b = rng.choice([b for b in built if b[3] > 1])
        b[3] -= 1
        total -= 1

    queries, answers, kinds = [], [], Counter()
    for name, kind, qs, k in built:
        for qtype, text, ans in rng.sample(qs, k):
            qid = len(queries) + 1
            queries.append({"id": qid, "file_name": name, "query": text})
            answers.append({"id": qid, "answer": ans})
            kinds[kind] += 1
    (out / "queries.json").write_text(json.dumps(queries, ensure_ascii=False, indent=2), encoding="utf-8")
    (out / "answers.json").write_text(json.dumps(answers, ensure_ascii=False, indent=2), encoding="utf-8")
    with open(out / "sample_submission.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["id", "answer"])
        for a in answers:
            w.writerow([a["id"], a["answer"]])

    print(f"{len(built)} file, {len(queries)} câu -> {out}")
    print("Số câu theo kiểu bảng:", dict(kinds))
    bad = {k[:2]: v for k, v in stats.items() if k[2] == "fail"}
    never = [k for k in {k[:2] for k in stats} if stats[k + ("ok",)] == 0]
    if never:
        print("CẢNH BÁO - loại câu không sinh được lần nào:", never)
    print("Số lần sinh lại theo loại câu:", {f"{a}.{b}": v for (a, b), v in sorted(bad.items())})


if __name__ == "__main__":
    main()
