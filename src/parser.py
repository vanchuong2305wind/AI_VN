from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

MISSING = {"", "none", "nan", "�", "-", "–", "—", "(na)", "n/a", "na", "...", "…", "(x)", "..", "(d)", "(s)",
           "(z)", "n.a.", "n.d."}
PLACEHOLDER_RE = re.compile(r"^\(-\d+-\)$|^-\d+-$"
                            r"|^\((?:[a-z]|[ivx]{2,4})\)$")
KEEP_HYPHEN = {"self", "non", "co", "full", "part", "year", "well", "high", "low", "long", "short", "multi"}
NOTE_RE = re.compile(r"^(footnotes?|notes?|source|sources)\b|^\(\d+\)|^\d+/\s", re.I)
FOOTNOTE_RE = re.compile(r"\s*\(\d+\)|\s+\d+/(?=\s|$)")
NUM_RE = re.compile(r"^[+-]?\(?\$?\d[\d,]*(\.\d+)?\)?%?$")
CODE_RE = re.compile(r"^(?=.*[A-Z])[A-Z0-9]{3,8}$")


@dataclass
class Table:
    id: str
    sheet: str
    kind: str
    df: pd.DataFrame
    columns: dict[str, str]
    context: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    label_header: str = ""
    long: pd.DataFrame | None = None
    labels: set[str] = field(default_factory=set)
    phys: dict[str, int] = field(default_factory=dict)
    code_header: bool = False


def clean_text(s: str) -> str:
    def join(m: re.Match) -> str:
        prefix = re.findall(r"[A-Za-z]+$", m.group(1))
        keep = prefix and prefix[0].lower() in KEEP_HYPHEN
        return f"{m.group(1)}-{m.group(2)}" if keep else m.group(1) + m.group(2)

    s = s.replace("\xa0", " ")
    s = re.sub(r"(\w+)-\s*\n\s*([a-z])", join, s)
    s = re.sub(r"\s+", " ", s).strip()
    return FOOTNOTE_RE.sub("", s).strip()


def clean_value(v):
    if v is None or isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return None if v != v else v
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, dt.date):
        return dt.datetime(v.year, v.month, v.day)
    s = str(v).replace("\xa0", " ").strip()
    if s.lower() in MISSING or PLACEHOLDER_RE.match(s):
        return None
    if NUM_RE.match(s):
        num = float(re.sub(r"[^\d.]", "", s))
        num = -num if s.startswith(("-", "(")) else num
        return int(num) if num.is_integer() and "." not in s else num
    return s


def is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def is_year(v) -> bool:
    return is_num(v) and float(v).is_integer() and 1900 <= v <= 2100


def to_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d")
    if is_num(v) and float(v).is_integer():
        return str(int(v))
    return clean_text(str(v))


def is_caps(s: str) -> bool:
    letters = re.sub(r"[^A-Za-z]", "", s)
    return len(letters) > 2 and letters.isupper()


class Grid:

    def __init__(self, ws):
        raw = {}
        self.indent = {}
        self.na = {}
        for row in ws.iter_rows():
            for c in row:
                value = getattr(c, "value", None)
                v = clean_value(value)
                if v is None and isinstance(value, str) and value.strip():
                    self.na[c.row] = self.na.get(c.row, 0) + 1
                if v is not None:
                    raw[(c.row, c.column)] = v
                    if isinstance(v, str):
                        lead = len(value) - len(value.lstrip(" \xa0　"))
                        align = getattr(getattr(c, "alignment", None), "indent", 0) or 0
                        if lead or align:
                            self.indent[(c.row, c.column)] = int(align) * 3 + lead
        self.nrows = max((r for r, _ in raw), default=0)
        self.ncols = max((c for _, c in raw), default=0)
        self.val = [[None] * (self.ncols + 1) for _ in range(self.nrows + 1)]
        for (r, c), v in raw.items():
            self.val[r][c] = v
        self.origin = {}
        merged = ws.merged_cells.ranges if hasattr(ws, "merged_cells") else []
        for rng in merged:
            r0, c0 = rng.min_row, rng.min_col
            if r0 > self.nrows or c0 > self.ncols:
                continue
            v = self.val[r0][c0]
            for r in range(r0, min(rng.max_row, self.nrows) + 1):
                for c in range(c0, min(rng.max_col, self.ncols) + 1):
                    self.val[r][c] = v
                    self.origin[(r, c)] = (r0, c0)

    def cells(self, r):
        out, seen = [], set()
        for c in range(1, self.ncols + 1):
            v = self.val[r][c]
            if v is None:
                continue
            key = self.origin.get((r, c), (r, c))[1]
            if key not in seen:
                seen.add(key)
                out.append((c, v))
        return out


def row_kind(grid: Grid, r: int, c0: int) -> str:
    cells = grid.cells(r)
    if not cells:
        return "empty"
    if (len(cells) == 2 and is_num(cells[0][1]) and float(cells[0][1]).is_integer() and 0 <= cells[0][1] <= 9
            and isinstance(cells[1][1], str)):
        return "single_label"
    nums = [v for _, v in cells if is_num(v) or isinstance(v, dt.datetime)]
    if nums and all(is_year(v) for v in nums):
        return "text" if grid.val[r][c0] is None else "years"
    na = grid.na.get(r, 0)
    if (nums or na) and len(nums) + na >= 0.4 * (len(cells) + na):
        return "data"
    if len(cells) == 1:
        return "single_label" if cells[0][0] == c0 else "single_span"
    return "text"


@dataclass
class Block:
    titles: list[int] = field(default_factory=list)
    header: list[int] = field(default_factory=list)
    body: list[tuple[int, str]] = field(default_factory=list)
    notes: list[int] = field(default_factory=list)


def split_blocks(grid: Grid, c0: int) -> list[Block]:
    kinds = {r: row_kind(grid, r, c0) for r in range(1, grid.nrows + 1)}
    blocks, buf, cur, in_notes = [], [], None, False

    def open_block(first_data_row: int) -> Block:
        rows = [r for r in buf if kinds[r] != "empty"]
        sections = []
        while rows and kinds[rows[-1]] == "single_label":
            sections.insert(0, rows.pop())
        header = []
        while rows and kinds[rows[-1]] in ("text", "single_span"):
            header.insert(0, rows.pop())
        return Block(titles=rows, header=header,
                     body=[(r, "section") for r in sections] + [(first_data_row, "data")])

    for r in range(1, grid.nrows + 1):
        k = kinds[r]
        if k == "years":
            k = kinds[r] = "text" if cur is None else "data"
        if cur is None:
            if k == "data":
                cur, buf = open_block(r), []
            else:
                buf.append(r)
            continue
        if k == "empty":
            continue
        text = str(grid.val[r][c0] or "")
        if k == "data" and not in_notes:
            cur.body.append((r, "data"))
        elif k in ("single_label", "data"):
            if in_notes or NOTE_RE.match(text) or len(text) > 120:
                in_notes = True
                cur.notes.append(r)
            else:
                cur.body.append((r, "section"))
        else:
            trailing = []
            while cur.body and cur.body[-1][1] == "section":
                trailing.insert(0, cur.body.pop()[0])
            blocks.append(cur)
            cur, buf, in_notes = None, trailing + [r], False
    if cur is not None:
        blocks.append(cur)
    return blocks


def detect_label_col(grid: Grid, data_rows: list[int], cols: list[int]):
    def text_ratio(c):
        vals = [grid.val[r][c] for r in data_rows if grid.val[r][c] is not None]
        return sum(isinstance(v, str) for v in vals) / len(vals) if vals else 0.0

    def filled(c):
        return sum(grid.val[r][c] is not None for r in data_rows)

    for c in cols[:3]:
        if text_ratio(c) >= 0.6:
            rest = [x for x in cols if x > c and filled(x) >= max(2, 0.1 * len(data_rows))] or [x for x in cols if x > c]
            return c if rest and all(text_ratio(x) < 0.3 for x in rest) else None
    return None


def section_paths(body, grid: Grid, label_col: int) -> dict[int, str]:
    path, had_data, out = [], True, {}
    for r, kind in body:
        if kind == "data":
            out[r] = " > ".join(path)
            had_data = True
            continue
        name = to_text(grid.val[r][label_col])
        if is_caps(name):
            path = [name]
        elif had_data:
            path = path[:-1] + [name]
        else:
            path = path + [name]
        had_data = False
    return out


def repeat_groups(labels: list[str]) -> list[str]:
    if len(set(labels)) == len(labels):
        return [""] * len(labels)
    groups, cur, seen = [""] * len(labels), labels[0], set()
    for i, lab in enumerate(labels):
        if lab in seen:
            cur = labels[i - 1]
            groups[i - 1] = cur
            seen = {cur}
        groups[i] = cur
        seen.add(lab)
    return groups


def join_header(parts: list[tuple[str, bool]]) -> str:
    out, prev_merged = "", True
    for text, merged in parts:
        wrapped = out and not merged and not prev_merged
        if wrapped and out.endswith("-") and text[:1].islower():
            word = re.findall(r"([A-Za-z]+)-$", out)
            out = out if word and word[0].lower() in KEEP_HYPHEN else out[:-1]
            out += text
        elif wrapped and (out[-1] in ",&/" or text[:1].islower() or text[:1] == "&"):
            out += " " + text
        else:
            out = f"{out} > {text}" if out else text
        prev_merged = merged
    return out


def indent_parents(df: pd.DataFrame, rows: list[int], grid: Grid, label_col: int) -> bool:
    levels = [grid.indent.get((r, label_col), 0) for r in rows]
    if len(set(levels)) < 2:
        return False
    parents, stack, section = [], [], None
    for lvl, lab, sec in zip(levels, df["label"], df["section"]):
        if sec != section:
            stack, section = [], sec
        while stack and stack[-1][0] >= lvl:
            stack.pop()
        parents.append(stack[-1][1] if stack else "")
        stack.append((lvl, lab))
    df["group"] = parents
    return True


def indent_groups(df: pd.DataFrame, columns: dict[str, str]) -> bool:
    ind = next((c for c, n in columns.items() if re.search(r"indent|level", n, re.I)), None)
    if ind is None or not pd.api.types.is_numeric_dtype(df[ind]):
        return False
    stack, parents = [], []
    for lvl, lab in zip(df[ind].fillna(0).astype(int), df["label"]):
        stack = stack[:lvl]
        parents.append(" > ".join(stack))
        stack = stack + [lab]
    df["group"] = parents
    return True


def build_table(grid: Grid, block: Block, c0: int, sheet: str, tid: str,
                named: dict[int, str] | None = None) -> Table | None:
    data_rows = [r for r, k in block.body if k == "data"]
    if not data_rows:
        return None
    cols = [c for c in range(1, grid.ncols + 1) if any(grid.val[r][c] is not None for r in data_rows)]
    label_col = detect_label_col(grid, data_rows, cols)
    value_cols = [c for c in cols if c != label_col]

    context = [to_text(v) for r in block.titles for _, v in grid.cells(r)]
    parts = {c: [] for c in value_cols}
    for h in block.header:
        texts = [to_text(grid.val[h][c]) for c in value_cols]
        if len(value_cols) > 1 and texts[0] and len(set(texts)) == 1:
            context.append(texts[0])
            continue
        for c, t in zip(value_cols, texts):
            if t and t not in (p for p, _ in parts[c]):
                parts[c].append((t, (h, c) in grid.origin))
    label_header = ""
    if label_col:
        label_header = " > ".join(dict.fromkeys(
            to_text(grid.val[h][label_col]) for h in block.header if grid.val[h][label_col] is not None))

    own = {c: join_header(parts[c]) for c in value_cols}
    texts = [p for c in value_cols for p, _ in parts[c]]
    code_header = (len(texts) >= 2 and all(CODE_RE.match(p) for p in texts)
                   and len({len(p) for p in texts}) == 1)
    inherit = named and sum(c in named for c in value_cols) >= 0.8 * len(value_cols) and (
        code_header or (not texts and set(value_cols) == set(named)))
    columns, seen, phys = {}, {}, {}
    for i, c in enumerate(value_cols, 1):
        name = own[c]
        if inherit and c in named:
            name = f"{named[c]} > {name}" if name else named[c]
        name = name or get_column_letter(c)
        seen[name] = seen.get(name, 0) + 1
        columns[f"c{i}"] = name if seen[name] == 1 else f"{name} #{seen[name]}"
        phys[f"c{i}"] = c

    sections = section_paths(block.body, grid, label_col) if label_col else {}
    records = []
    for r in data_rows:
        rec = {"row": r, "section": sections.get(r, ""), "group": ""}
        if label_col:
            rec["label"] = to_text(grid.val[r][label_col])
        for i, c in enumerate(value_cols, 1):
            v = grid.val[r][c]
            rec[f"c{i}"] = clean_text(v) if isinstance(v, str) else v
        records.append(rec)
    df = pd.DataFrame(records)
    for c in columns:
        if df[c].map(lambda v: v is None or is_num(v)).all():
            df[c] = pd.to_numeric(df[c], errors="coerce")
        elif df[c].map(lambda v: v is None or isinstance(v, dt.datetime)).all():
            df[c] = pd.to_datetime(df[c])

    kind = "matrix" if label_col else "records"
    if kind == "matrix" and not indent_groups(df, columns) and not indent_parents(df, data_rows, grid, label_col):
        for _, idx in df.groupby("section", sort=False).groups.items():
            df.loc[idx, "group"] = repeat_groups(df.loc[idx, "label"].tolist())

    t = Table(id=tid, sheet=sheet, kind=kind, df=df, columns=columns, label_header=label_header,
              context=[x for x in dict.fromkeys(context) if x],
              notes=[to_text(grid.val[r][c0]) for r in block.notes], phys=phys, code_header=code_header)
    t.long = to_long(t)
    t.labels = (set(columns.values()) | set(df["section"]) | set(df["group"])
                | {p for v in columns.values() for p in v.split(" > ")}
                | (set(df["label"]) if "label" in df else set())
                | {v for c in columns if df[c].dtype == object for v in df[c].dropna() if isinstance(v, str)})
    t.labels.discard("")
    return t


def to_long(t: Table) -> pd.DataFrame | None:
    if t.kind != "matrix":
        return None
    num_cols = [c for c in t.columns if pd.api.types.is_numeric_dtype(t.df[c])]
    long = t.df.melt(id_vars=["row", "section", "group", "label"], value_vars=num_cols,
                     var_name="col", value_name="value").dropna(subset=["value"])
    long.insert(5, "col_name", long["col"].map(t.columns))
    return long.reset_index(drop=True)


def parse_workbook(path: str | Path) -> list[Table]:
    path = Path(path)
    big = path.stat().st_size > 20 * 1024 * 1024
    wb = openpyxl.load_workbook(path, data_only=True, read_only=big)
    tables = []
    for ws in wb.worksheets:
        grid = Grid(ws)
        if grid.nrows == 0:
            continue
        c0 = min(c for c in range(1, grid.ncols + 1)
                 if any(grid.val[r][c] is not None for r in range(1, grid.nrows + 1)))
        named = {}
        for block in split_blocks(grid, c0):
            t = build_table(grid, block, c0, ws.title, f"T{len(tables) + 1}", named)
            if t is not None:
                tables.append(t)
                if not t.code_header:
                    named = {t.phys[c]: name for c, name in t.columns.items()}
    wb.close()
    return tables
